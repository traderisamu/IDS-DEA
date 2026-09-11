"""
Generates DEA_Config.xlsx - the admin-editable configuration file.
Run this once to (re)create the config file from the source data,
then place it on the shared network path. Admin can hand-edit the
resulting file directly in Excel afterwards (add/remove Work
Descriptions, tooltips, or employees) - no code changes needed.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# (Work Description, Details Tooltip, Restricted To JOB Codes, Exclusive).
# The 3rd value is a comma-separated list of JOB code PREFIXES (the part
# before the first "-", e.g. "ASPM" out of "ASPM-CUB") - leave it "" for
# an item that should always be available regardless of JOB code. The
# 4th value, "Yes" or "", makes a restricted item EXCLUSIVE: for any of
# its listed prefixes, this item becomes the ONLY Work Description shown
# - every normally-always-available item is hidden too, not just other
# restricted ones. Used for JOB codes dedicated to one specific activity
# (e.g. a meeting- or training-only code) that shouldn't be usable for
# anything else. Plain (unmarked) items are listed first, alphabetically;
# items with a leading "*"/"~" marker (a purely cosmetic grouping cue -
# see the README sheet) are listed after, also alphabetically by the
# visible word ignoring that marker.
WORK_ITEMS = [
    ("Assemble Calc Package", "", "", ""),
    ("Create Calculation Template", "Calc Name needs to be consistent", "", ""),
    ("Create IER/Connection Scheme", "", "", ""),
    ("Create RFI or Sketches", "RFI#", "", ""),
    ("Review Shop & Erection Dwgs", "# sheets", "", ""),
    ("Run Calculation Joints", "# of joints", "", ""),
    ("Update Calculation Runs", "# of joints", "", ""),
    ("Update Calculation Templates", "# of calcs", "", ""),
    ("*Check Calculation Runs", "# of joints", "", ""),
    ("*Check Calculation Templates", "Calc Name", "", ""),
    ("*Meeting", "", "ASPM,AMGM", "Yes"),
    ("~Read/Respond Emails", "", "", ""),
    ("*Reading References", "", "ENGG,AESG", ""),
    ("~Review CD/RFI Response", "", "", ""),
    ("*Training", "", "ASPT", "Yes"),
]

# (Employee Name, Team, Employee ID). Employee ID is "MM-YYYY-####" (hire
# month/year + employee number). The trailing number IS that person's
# PIN, read live by the app every time it starts - there's no separate
# PIN storage and no seeding step. To set, change, or remove someone's
# PIN, just edit their Employee ID here and save.
EMPLOYEES = [
    ("Abegail Sena", "Miko", "06-2010-526"),
    ("Oscar Tolledo Jr.", "Oscar", "10-2011-678"),
    ("Angelo Rafael Agbayani", "Angelo", "04-2013-875"),
    ("Abegail Lara Oaña", "Lara", "12-2013-906"),
    ("Mikael Oliver Umali", "Miko", "03-2014-920"),
    ("Cammille Estrella", "Miko", "05-2014-924"),
    ("Al Capunay", "Fernando", "03-2015-973"),
    ("Joan Villanueva", "Fernando", "07-2016-1020"),
    ("Jay Benedict Cabrera", "Miko", "02-2018-1096"),
    ("John Luer Suaiso", "Luer", "07-2018-1114"),
    ("Jericho Fiestada", "Oscar", "08-2018-1119"),
    ("Santana Gozon", "Luer", "07-2019-1170"),
    ("Maribel Rubis", "Lara", "04-2021-1199"),
    ("Rodel Marticio", "Miko", "02-2022-1223"),
    ("Anthony Manigo", "Angelo", "02-2022-1221"),
    ("Mary Joyce Dela Cruz", "Angelo", "06-2022-1229"),
    ("Richard Manalises", "Luer", "05-2022-1228"),
    ("Colene Gammad", "Miko", "07-2023-1269"),
    ("Ediel Pagani", "Lara", "06-2023-1262"),
    ("Lars Christian Villena", "Luer", "02-2024-1264"),
    ("Andrew Bernardino", "Miko", "05-2024-1319"),
    ("Joseph Melo Manalo", "Luer", "10-2024-1330"),
    ("Angelo Carlo Lazaro", "Angelo", "06-2025-1347"),
    ("Aldrin Jonas Eduarte", "Luer", "06-2025-1348"),
    ("Richelle Anne Torres", "Fernando", "04-2026-1381"),
    ("Ivan Randel Fontillas", "Oscar", "04-2026-1382"),
    ("Debbie Kaye Marie Caballero", "Fernando", "06-2026-1383"),
    ("Ralp Gumiling", "Lara", "06-2026-1384"),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="1F4E78")
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="808080")
BODY_FONT = Font(name="Arial", size=11)
YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(out_path=None):
    if out_path is None:
        out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "DEA_Config.xlsx")
    wb = openpyxl.Workbook()

    # ---------- README ----------
    ws = wb.active
    ws.title = "READ ME"
    ws["A1"] = "IDS PH - Daily Employee Accomplishment (DEA) Logger - Config"
    ws["A1"].font = TITLE_FONT
    lines = [
        "",
        "This file controls the dropdown lists and tooltips used by the DEA Logger app.",
        "Edit it directly in Excel and save - every user's app re-reads it automatically",
        "(cached for a few minutes, or they can click 'Refresh Config' in the app).",
        "",
        "Sheets:",
        "  - 'Work Items'  : the Work Description dropdown + the tooltip shown under Details.",
        "                    Add a row to add a new option. Do not leave the Description blank.",
        "                    The 3rd column optionally restricts an item to only show when the",
        "                    JOB Code being typed starts with one of the listed prefixes (comma-",
        "                    separated, e.g. 'ASPM,AMGM') - leave it blank for an item that should",
        "                    always be available. The 4th column, 'Exclusive', set to Yes, makes a",
        "                    restricted item the ONLY one shown for its listed prefixes - every",
        "                    normally-always-available item is hidden too for that JOB Code, not",
        "                    just other restricted ones. Use this for a JOB Code dedicated to one",
        "                    activity only (e.g. a meeting- or training-only code). A leading '*'",
        "                    or '~' in the Description itself is purely cosmetic (a visual grouping",
        "                    cue) - it does nothing on its own; the JOB-code behavior only comes",
        "                    from the 3rd/4th columns.",
        "                    The JOB Code field has no dropdown list of its own - it's free-text,",
        "                    auto-suggesting each employee's own previously-used codes.",
        "  - 'Employees'   : the full employee roster, Team, and Employee ID.",
        "                    Add/remove rows to add or remove people from the app.",
        "                    The trailing number of Employee ID IS that person's PIN",
        "                    (see PINs note below) - it's read live, not just at setup.",
        "  - 'App Settings': reminder times, hour thresholds, admin login, shared paths.",
        "",
        "Do not rename the sheet tabs or column headers - the app looks them up by name.",
        "",
        "PINs: there is no separate PIN storage. Each employee's PIN is simply the",
        "number at the end of their Employee ID above (e.g. ID 06-2010-526 -> PIN 526),",
        "read fresh from this file every time the app starts. Employees cannot set,",
        "change, or view their own PIN from the app - to set, change, or remove",
        "someone's PIN, edit their Employee ID here and save (or Refresh Config in the",
        "Admin Dashboard, which also has a read-only 'View Employee PINs...' lookup).",
    ]
    for i, line in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=line).font = BODY_FONT if line and not line.startswith(" " * 2) else NOTE_FONT
    autosize(ws, [100])

    # ---------- Work Items ----------
    ws = wb.create_sheet("Work Items")
    ws["A1"] = "Work Description"
    ws["B1"] = "Details Tooltip (shown to user under the Details field)"
    ws["C1"] = "Restricted To JOB Codes (optional, comma-separated prefixes)"
    ws["D1"] = "Exclusive (Yes = ONLY this item shows for those JOB codes)"
    style_header(ws, 1, 4)
    for i, (desc, tip, prefixes, exclusive) in enumerate(WORK_ITEMS, start=2):
        ws.cell(row=i, column=1, value=desc).font = BODY_FONT
        ws.cell(row=i, column=2, value=tip).font = BODY_FONT
        ws.cell(row=i, column=3, value=prefixes).font = BODY_FONT
        ws.cell(row=i, column=4, value=exclusive).font = BODY_FONT
    autosize(ws, [38, 55, 42, 30])
    ws.freeze_panes = "A2"

    # ---------- Employees ----------
    ws = wb.create_sheet("Employees")
    ws["A1"] = "Employee Name"
    ws["B1"] = "Team"
    ws["C1"] = "Employee ID"
    style_header(ws, 1, 3)
    for i, (name, team, emp_id) in enumerate(EMPLOYEES, start=2):
        ws.cell(row=i, column=1, value=name).font = BODY_FONT
        ws.cell(row=i, column=2, value=team).font = BODY_FONT
        ws.cell(row=i, column=3, value=emp_id).font = BODY_FONT
    autosize(ws, [32, 16, 16])
    ws.freeze_panes = "A2"

    # ---------- App Settings ----------
    ws = wb.create_sheet("App Settings")
    ws["A1"] = "Setting"
    ws["B1"] = "Value"
    ws["C1"] = "Notes"
    style_header(ws, 1, 3)
    settings = [
        ("CompanyName", "IDS PH", "Shown in the app title bar / header."),
        ("SharedDataPath",
         r"\\EgnyteDrive\idsinc\Shared\Engineering\ENGG PHL\REPORTS\16 Manhour Report LEADERBOARD\DEA App",
         "Network folder holding this config file AND the Logs\\ subfolder (one Excel per employee). Edit the app's local pointer in dea_app_settings.json if this ever moves - see README."),
        ("ReminderTimes", "09:00,15:00",
         "Comma-separated 24h HH:MM times, earliest first. The EARLIEST time (e.g. 9:00 AM) "
         "checks whether the PREVIOUS work day's log was completed and pops up once if not - "
         "it does not repeat. Every time AFTER that (e.g. 3:00 PM) checks TODAY's log instead; "
         "once the last of those has passed with today still incomplete, it keeps re-nagging "
         "every ReminderRepeatMinutes until ReminderRepeatCutoffMinutes runs out, or until "
         "something is logged, or the day is marked Holiday/Leave."),
        ("ReminderRepeatMinutes", "30",
         "How often to keep re-nagging, in minutes, after the LAST time in ReminderTimes has "
         "passed with today's log still incomplete. Floored at 15 (never fires more often than "
         "every 15 minutes) no matter how low this is set, so a typo here can't turn into a "
         "disruptive rapid-fire popup."),
        ("ReminderRepeatCutoffMinutes", "90",
         "How long, in minutes, to keep re-nagging after the LAST time in ReminderTimes before "
         "giving up on that day entirely - even if still incomplete. Default 90 with the default "
         "3:00 PM slot and 30-minute repeat means popups at 3:00, 3:30, 4:00 and 4:30, then "
         "nothing more that day."),
        ("MinHoursForGreenDay", 4, "Calendar day only turns GREEN if total logged hours >= this value."),
        ("MaxHoursWarningThreshold", 8, "App shows a warning if total hours logged for one day exceeds this value."),
        ("AdminUsername", "admin", "Login for the Admin Dashboard (all-employee calendar, DB folder link)."),
        ("AdminPassword", "nimda", "Change this, then tell your team to update their config path only if you move the file - the password is re-read live from here."),
        ("WeekendsCountAsWorkday", "No", "If No, Saturdays/Sundays are shown neutral (not red) even with no log."),
    ]
    for i, (k, v, note) in enumerate(settings, start=2):
        ws.cell(row=i, column=1, value=k).font = BODY_FONT
        vcell = ws.cell(row=i, column=2, value=v)
        vcell.font = BODY_FONT
        vcell.fill = YELLOW_FILL
        ws.cell(row=i, column=3, value=note).font = NOTE_FONT
    autosize(ws, [26, 70, 70])
    ws.freeze_panes = "A2"

    wb.save(out_path)
    print(f"Saved {out_path}")


if __name__ == "__main__":
    build()
