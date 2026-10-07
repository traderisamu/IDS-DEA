"""
Reads DEA_Config.xlsx from the shared network path.

This is the ONLY place hardcoded business data (job/work-description
list, tooltips, employee roster, reminder time, hour thresholds, admin
login) is allowed to come from. The admin edits the Excel file; nothing
here should ever need a code change for a roster/list update.
"""
import os
import re
import time
import threading
import openpyxl

from . import constants as C


class ConfigError(Exception):
    pass


# OLE compound-document magic: a password-encrypted .xlsx is an OLE
# container, while a plain .xlsx is a ZIP ("PK.."). Lets us tell
# "encrypted, needs the password" apart from "corrupt/unreadable".
_OLE_SIGNATURE = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"


def _looks_encrypted(path):
    try:
        with open(path, "rb") as f:
            return f.read(8) == _OLE_SIGNATURE
    except OSError:
        return False


def open_config_workbook(path):
    """Opens DEA_Config.xlsx whether it is plain or file-open-encrypted.

    Plain files load directly (pre-encryption rollout, local seed copies).
    Encrypted files are decrypted in memory with the bundled
    CONFIG_FILE_PASSWORD - no plaintext copy ever touches disk. Raises
    ConfigError with an actionable message instead of a raw traceback
    when decryption fails (old build, or password rotated without a
    matching app update)."""
    try:
        return openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        if not os.path.isfile(path) or not _looks_encrypted(path):
            raise
    try:
        import io
        import msoffcrypto
        buf = io.BytesIO()
        with open(path, "rb") as f:
            office_file = msoffcrypto.OfficeFile(f)
            office_file.load_key(password=C.CONFIG_FILE_PASSWORD)
            office_file.decrypt(buf)
        buf.seek(0)
        wb = openpyxl.load_workbook(buf, data_only=True, read_only=True)
        # read_only streams lazily from `buf` - keep it alive with the book.
        wb._decrypted_config_buf = buf
        return wb
    except ConfigError:
        raise
    except Exception as e:
        raise ConfigError(
            "Could not open the password-protected config file - this app "
            "build could not decrypt it. Update to the latest build; if it "
            "still fails, the config password was changed without shipping "
            "a matching app update."
        ) from e


class AppConfig:
    """Holds one snapshot of the config workbook contents."""

    def __init__(self):
        self.work_items = []          # list[(description, tooltip)]
        self.work_item_prefixes = {}  # dict[description -> frozenset(JOB prefixes)], only for restricted items
        self.work_item_exclusive = set()  # descriptions marked "Exclusive" - see work_items_for_job
        self.employees = []           # list[(name, team)]
        self.employee_ids = {}        # dict[name -> Employee ID string], optional column
        self.settings = {}            # dict[str, str]
        self.loaded_at = 0.0
        self.source_path = None

    # -- convenience accessors with fallback defaults -----------------
    def reminder_times(self):
        """Returns a sorted list of 'HH:MM' strings. Falls back to the
        single legacy ReminderTime setting, then to the hardcoded default."""
        raw = self.settings.get("ReminderTimes")
        if not raw:
            raw = self.settings.get("ReminderTime", C.DEFAULT_REMINDER_TIME)
        times = []
        for part in str(raw).split(","):
            part = part.strip()
            if not part:
                continue
            try:
                hh, mm = part.split(":")
                hh, mm = int(hh), int(mm)
                if 0 <= hh <= 23 and 0 <= mm <= 59:
                    times.append(f"{hh:02d}:{mm:02d}")
            except (ValueError, TypeError):
                continue
        return sorted(times) or [C.DEFAULT_REMINDER_TIME]

    def reminder_repeat_minutes(self):
        try:
            val = int(float(self.settings.get(
                "ReminderRepeatMinutes", C.DEFAULT_REMINDER_REPEAT_MINUTES)))
        except (TypeError, ValueError):
            val = C.DEFAULT_REMINDER_REPEAT_MINUTES
        return max(val, C.MIN_REMINDER_REPEAT_MINUTES)

    def reminder_repeat_cutoff_minutes(self):
        """How long after the last 'today' reminder slot to keep
        repeating before stopping for the day entirely, no matter what."""
        try:
            val = int(float(self.settings.get(
                "ReminderRepeatCutoffMinutes", C.DEFAULT_REMINDER_REPEAT_CUTOFF_MINUTES)))
        except (TypeError, ValueError):
            val = C.DEFAULT_REMINDER_REPEAT_CUTOFF_MINUTES
        return max(val, self.reminder_repeat_minutes())  # always room for at least one repeat

    def min_hours_green(self):
        try:
            return float(self.settings.get(
                "MinHoursForGreenDay", C.DEFAULT_MIN_HOURS_GREEN))
        except (TypeError, ValueError):
            return C.DEFAULT_MIN_HOURS_GREEN

    def max_hours_warning(self):
        try:
            return float(self.settings.get(
                "MaxHoursWarningThreshold", C.DEFAULT_MAX_HOURS_WARNING))
        except (TypeError, ValueError):
            return C.DEFAULT_MAX_HOURS_WARNING

    def admin_user(self):
        return str(self.settings.get("AdminUsername", C.DEFAULT_ADMIN_USER))

    def admin_pass(self):
        return str(self.settings.get("AdminPassword", C.DEFAULT_ADMIN_PASS))

    def company_name(self):
        return self.settings.get("CompanyName", C.COMPANY_NAME)

    def weekends_count_as_workday(self):
        val = str(self.settings.get("WeekendsCountAsWorkday", "No")).strip().lower()
        return val in ("yes", "y", "true", "1")

    def team_of(self, employee_name):
        for name, team in self.employees:
            if name == employee_name:
                return team
        return None

    def tooltip_for(self, work_description):
        for desc, tip in self.work_items:
            if desc == work_description:
                return tip
        return ""

    def work_items_for_job(self, job_code, always_include=None):
        """Work Descriptions available for the given JOB code, in the
        same order as the 'Work Items' sheet.

        Most items have no restriction and always show up. A few (see
        the 'Restricted To JOB Codes' column) only show when the JOB
        code's prefix - everything before the first '-', e.g. 'ASPM'
        out of 'ASPM-CUB' - matches one of the prefixes listed for that
        item; those are ADDED on top of the normal always-available
        ones. Matching is case-insensitive and ignores whitespace.

        A restricted item can also be marked 'Exclusive' (see that
        column) - if the JOB code's prefix matches an EXCLUSIVE item,
        that item (and any other item ALSO marked exclusive for that
        same prefix) is the ONLY thing shown - every normal
        always-available item, and any other item's restriction, is
        hidden for that prefix. This is for JOB codes that are
        dedicated entirely to one specific activity (e.g. a meeting- or
        training-only JOB code) and shouldn't be usable for anything
        else.

        If a JOB code doesn't match any prefix at all (including a
        blank/not-yet-typed JOB code), restricted items simply don't
        show - not an error, just not relevant yet.

        `always_include`, if given, is a Work Description that's kept in
        the returned list even if it wouldn't otherwise qualify - used
        when editing an existing entry so its current value doesn't
        vanish out from under the dropdown just because the JOB code was
        since changed or the rule was tightened."""
        prefix = str(job_code or "").strip().split("-")[0].strip().upper()

        exclusive_matches = [
            desc for desc, _tip in self.work_items
            if desc in self.work_item_exclusive and prefix and prefix in self.work_item_prefixes.get(desc, frozenset())
        ]
        if exclusive_matches:
            if always_include and always_include not in exclusive_matches:
                exclusive_matches.append(always_include)
            return exclusive_matches

        result = []
        for desc, _tip in self.work_items:
            allowed = self.work_item_prefixes.get(desc)
            if not allowed or prefix in allowed or desc == always_include:
                result.append(desc)
        return result

    def employee_id_for(self, employee_name):
        """Returns the 'Employee ID' column value for this person, or
        None if that (optional) column is missing or blank for them."""
        return self.employee_ids.get(employee_name)

    def pin_for(self, employee_name):
        """The employee's PIN, straight from DEA_Config.xlsx - there is
        no separately stored/settable PIN anymore. It's the trailing run
        of digits in their 'Employee ID' column (e.g. Employee ID
        '06-2010-526' -> PIN '526'). Returns None if that column is
        blank or missing for them, in which case the app just doesn't
        lock for that person - editing the Employee ID in the Employees
        sheet is the only way to set, change, or remove someone's PIN."""
        emp_id = self.employee_id_for(employee_name)
        if not emp_id:
            return None
        match = re.search(r"(\d+)\s*$", str(emp_id))
        return match.group(1) if match else None


def _read_config_file(path):
    if not os.path.isfile(path):
        raise ConfigError(f"Config file not found:\n{path}")

    wb = open_config_workbook(path)
    cfg = AppConfig()
    cfg.source_path = path

    # Work Items
    if "Work Items" not in wb.sheetnames:
        raise ConfigError("Config file is missing the 'Work Items' sheet.")
    ws = wb["Work Items"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        desc = row[0]
        tip = row[1] if len(row) > 1 else ""
        prefixes_raw = row[2] if len(row) > 2 else None
        exclusive_raw = row[3] if len(row) > 3 else None
        if desc:
            desc = str(desc).strip()
            cfg.work_items.append((desc, str(tip).strip() if tip else ""))
            if prefixes_raw and str(prefixes_raw).strip():
                prefixes = {p.strip().upper() for p in re.split(r"[,/]", str(prefixes_raw)) if p.strip()}
                if prefixes:
                    cfg.work_item_prefixes[desc] = frozenset(prefixes)
            if str(exclusive_raw or "").strip().lower() in ("yes", "y", "true", "1"):
                cfg.work_item_exclusive.add(desc)

    # Employees
    if "Employees" not in wb.sheetnames:
        raise ConfigError("Config file is missing the 'Employees' sheet.")
    ws = wb["Employees"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        name = row[0]
        team = row[1] if len(row) > 1 else ""
        emp_id = row[2] if len(row) > 2 else None
        if name:
            name = str(name).strip()
            cfg.employees.append((name, str(team).strip() if team else ""))
            if emp_id:
                cfg.employee_ids[name] = str(emp_id).strip()

    # App Settings
    if "App Settings" in wb.sheetnames:
        ws = wb["App Settings"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            key = row[0]
            val = row[1] if len(row) > 1 else None
            if key:
                cfg.settings[str(key).strip()] = val

    cfg.loaded_at = time.time()
    wb.close()
    return cfg


def load_config_with_fallback(candidate_dirs, filename, attempts=4, delay_seconds=0.5):
    """
    Tries each folder in `candidate_dirs`, in order, looking for
    `filename`. For each candidate, retries a few times with a short
    delay before moving on to the next one.

    Why retry at all: the most common cause of "it says the config file
    is missing even though it's right there" is the app auto-starting at
    Windows login *before* the network/mapped drive holding the shared
    folder has finished mounting - a fraction-of-a-second-to-few-seconds
    race that has nothing to do with the path being wrong. A blind retry
    with a short pause resolves that without bothering anyone.

    Why multiple candidates: supports both the normal shared-network-
    drive deployment and a simpler "config file copied next to the exe"
    local/portable deployment, without the user ever having to type a
    path for the common cases.

    Returns (AppConfig, resolved_folder) on success - resolved_folder is
    whichever candidate worked, so the caller can remember it for next
    time. Raises ConfigError, with the specific problem found at every
    candidate, if all of them fail.
    """
    tried = []
    seen_dirs = set()
    for folder in candidate_dirs:
        if not folder or folder in seen_dirs:
            continue
        seen_dirs.add(folder)
        path = os.path.join(folder, filename)
        last_err = None
        for attempt in range(attempts):
            try:
                cfg = _read_config_file(path)
                return cfg, folder
            except ConfigError as e:
                last_err = e
                if attempt < attempts - 1:
                    time.sleep(delay_seconds)
        tried.append(f"  - {path}\n    {last_err}")

    detail = "\n".join(tried) if tried else "  (no candidate locations to try)"
    raise ConfigError(
        "Could not load the config file from any known location:\n\n" + detail
    )


class ConfigManager:
    """
    Thread-safe cached loader. Call get() everywhere in the UI; it
    refreshes from disk automatically every `ttl_seconds`, and callers
    can force an immediate refresh (e.g. an explicit "Refresh Config"
    button).
    """

    def __init__(self, config_path, ttl_seconds=180):
        self._path = config_path
        self._ttl = ttl_seconds
        self._lock = threading.Lock()
        self._cfg = None

    @property
    def path(self):
        return self._path

    @path.setter
    def path(self, new_path):
        with self._lock:
            self._path = new_path
            self._cfg = None  # force a fresh read from the new location

    def get(self, force_refresh=False):
        with self._lock:
            stale = (
                self._cfg is None
                or force_refresh
                or (time.time() - self._cfg.loaded_at) > self._ttl
            )
            if stale:
                self._cfg = _read_config_file(self._path)
            return self._cfg
