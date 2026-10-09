"""NaviTool 2.0, embedded as a tab inside the DEA Logger.

Originally a standalone app (Project - Job-nav ChatGPT v2); now imported
by client/main_app.py and mounted as the "NaviTool 2.0" notebook tab.
Must stay import-safe: no Tk root creation and no demo runner at import.
"""
import json
import os
import re
import subprocess
import sys
import threading
import time
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    HAS_DND = True
except Exception:
    HAS_DND = False

APP_NAME = "NaviTool 2.0"
JOB_ROOT = r"\\EgnyteDrive\idsinc\Shared\Engineering\ENGG PHL\JOBS"

# Work-stage groups. Each section header is a click-to-collapse toggle.
# Items are either plain buttons or dropdown (Menubutton) buttons; "{code}"
# in a relative path is replaced with the job's 3-letter code. Only items
# whose folder actually exists for the selected job are shown.
STAGE_TITLES = ["CALCS", "INITIAL", "SUBMITTAL", "FABRICATION"]

_CODE_RE = re.compile(r"\b([A-Za-z]{3})-")


def parse_job_code(folder_name):
    """'26 SPS-ASD15 Trinity' -> 'SPS'. None if the name has no CODE- pattern."""
    m = _CODE_RE.search(folder_name)
    return m.group(1).upper() if m else None


def _exists(base, rel):
    return bool(rel) and os.path.isdir(os.path.join(base, rel))


def _folder_menu(base, rel, depth=0, max_depth=3, label=None):
    """Nested menu entries for one folder: [(label, rel-or-children), ...].

    Every level starts with the folder's own name as an open-folder
    command (tk cascade headers can't run commands themselves, so this
    keeps each folder clickable: BS, BS > CALCS, BS > CALCS > BS01...),
    followed by one cascade per subfolder, down to max_depth. A folder
    with no subfolders of its own collapses to a plain open-folder
    command (a bare rel string)."""
    own = label if label is not None else (os.path.basename(rel.rstrip("\\")) or rel)
    entries = [(own, rel)]
    if depth >= max_depth:
        return entries
    kids = _dir_entries(base, rel)
    if kids:
        entries.append(("-", None))
        for name in kids:
            child = rel + "\\" + name if rel else name
            sub = _folder_menu(base, child, depth + 1, max_depth)
            entries.append((name, sub if len(sub) > 1 else child))
    return entries


def _dir_entries(base, rel):
    """Sorted list of immediate subfolders under a job-relative path."""
    out = []
    try:
        with os.scandir(os.path.join(base, rel)) as it:
            for e in it:
                if e.is_dir():
                    out.append(e.name)
    except OSError:
        pass
    return sorted(out, key=str.lower)


def _cascade_or_open(base, rel, label=None):
    """_folder_menu() when the folder has children, else a bare rel that
    renders as a plain open-folder command."""
    if _dir_entries(base, rel):
        return _folder_menu(base, rel, label=label)
    return rel


# A cascade level holding more than this many entries gets split into
# balanced "first to last" sub-cascades, so no native menu grows taller
# than the screen. Short levels render exactly as before.
MENU_CHUNK_SIZE = 25


def _chunk_range_label(first, last):
    """Compact "CD#001 to CD#025" / "MC01 to MC35" / "031226 to 033126"
    chunk label from the chunk's real first/last entry names. Two rules,
    in order: (1) same leading text plus a trailing number on both sides
    collapses to first-full-name + last-number (number strings kept
    verbatim, so zero-padding and "#" styles survive); (2) same embedded
    PREFIX#/-digits code with different numbers (the CD number buried in
    "262601 FBD CD#001 (Sequence Map...)") collapses to "CD#001 to
    CD#025". Single-letter codes (W21) and digit-free names can never
    match. Anything else falls back to the full "first to last" form,
    truncated so one long folder name can't blow out menu width."""
    first, last = str(first or "").strip(), str(last or "").strip()
    m1 = re.match(r"^(.*?)(\d+)$", first)
    m2 = re.match(r"^(.*?)(\d+)$", last)
    if m1 and m2 and m1.group(1) == m2.group(1) and m1.group(2) != m2.group(2):
        return "{} to {}{}".format(first[:28], m1.group(1), m2.group(2))[:60]
    c1 = re.search(r"([A-Z]{2,}[#-]?)(\d{1,4})", first)
    c2 = re.search(r"([A-Z]{2,}[#-]?)(\d{1,4})", last)
    if c1 and c2 and c1.group(1) == c2.group(1) and c1.group(2) != c2.group(2):
        return "{} to {}{}".format(c1.group(0), c1.group(1), c2.group(2))
    label = "{} to {}".format(first, last)
    return label if len(label) <= 60 else label[:57] + "..."


def _chunk_menu_entries(entries):
    """Split long runs of non-separator entries into balanced sub-cascades
    ("CD#1 to CD#60", ...). Leaves and sub-cascades chunk together (CD
    folders and CALCS connections mix both); separators, head pins, and
    short levels pass through untouched, in original order. Applied per
    level by _fill_menu, so every depth chunks independently."""
    out = []
    run = []

    def _flush():
        if len(run) > MENU_CHUNK_SIZE:
            n = -(-len(run) // MENU_CHUNK_SIZE)  # ceil: balanced, no tiny tail
            size, extra = divmod(len(run), n)
            idx = 0
            for c in range(n):
                take = size + (1 if c < extra else 0)
                part = run[idx:idx + take]
                idx += take
                out.append((_chunk_range_label(part[0][0], part[-1][0]),
                            list(part)))
        else:
            out.extend(run)
        del run[:]

    for entry in entries or []:
        if not entry or entry[0] == "-":
            _flush()
            if entry:
                out.append(entry)
        else:
            run.append(entry)
    _flush()
    return out


def build_job_sections(job_name, progress=None, cancel=None):
    """Return [(title, [item, ...]), ...]. An item is either
    {"kind":"button", label, rel, icon} or
    {"kind":"dropdown", label, icon, menu:[(menu_label, rel-or-children), ...]}
    where a nested list value is a further cascade (see _folder_menu).

    progress(phase_text) reports where the background build is (same
    contract as the Latest Details scan); cancel() returning True aborts
    with None (the caller keeps its placeholder and caches nothing)."""
    def _report(text):
        if progress:
            try:
                progress(text)
            except Exception:
                pass

    def _aborted():
        try:
            return bool(cancel and cancel())
        except Exception:
            return False

    base = os.path.join(JOB_ROOT, job_name)
    code = parse_job_code(job_name)
    if not code:
        return [("", [{"kind": "button", "label": "Job folder", "icon": "\U0001F4C1",
                       "rel": "", "menu": None}])]

    def sc(name, *parts):
        return "\\".join([code + "_" + name] + list(parts))

    sections = []

    # ---------- CALCS (main work) ----------
    calcs = []
    mathcad = sc("CALCS", "MATHCAD CALCS")
    if _exists(base, mathcad):
        menu = [("CALCS", sc("CALCS")), ("-", None)]
        for sub in ("MISC", "STRUCTURAL"):
            if _aborted():
                return None
            root = mathcad + "\\" + sub
            if _exists(base, root):
                submenu = [(sub, root), ("-", None)]
                names = _dir_entries(base, root)
                for i, name in enumerate(names):
                    if _aborted():
                        return None
                    conn = root + "\\" + name
                    submenu.append((name, _folder_menu(base, conn)))
                    _report("Reading CALCS… {}/{}".format(i + 1, len(names)))
                menu.append((sub, submenu))
        calcs.append({"kind": "dropdown", "label": "CALCS", "icon": "\U0001F9EE", "menu": menu})
    elif _exists(base, sc("CALCS")):
        calcs.append({"kind": "button", "label": "CALCS", "icon": "\U0001F9EE",
                      "rel": sc("CALCS"), "menu": None})
    if _aborted():
        return None
    sent = sc("CALCS", "SENT CALCS")
    if _exists(base, sent):
        menu = [("SENT CALCS", sent), ("-", None)]
        names = _dir_entries(base, sent)
        for i, name in enumerate(names):
            if _aborted():
                return None
            menu.append((name, sent + "\\" + name))
            _report("Reading SENT CALCS… {}/{}".format(i + 1, len(names)))
        calcs.append({"kind": "dropdown", "label": "SENT CALCS", "icon": "\U0001F4E4", "menu": menu})
    if calcs:
        sections.append(("CALCS", calcs))

    # ---------- Initial proceedings ----------
    if _aborted():
        return None
    _report("Reading INITIAL…")
    initial = []
    dd_root = sc("DESIGN DRAWINGS")
    if _exists(base, dd_root):
        menu = [("DESIGN DRAWINGS", dd_root), ("-", None)]
        for name in ("ARCHITECTURAL", "STRUCTURAL"):
            if _exists(base, dd_root + "\\" + name):
                menu.append((name, dd_root + "\\" + name))
        initial.append({"kind": "dropdown", "label": "DESIGN DRAWINGS",
                        "icon": "\U0001F4D0", "menu": menu})
    ier_root = sc("INITIAL ENGINEERING REVIEW")
    if _exists(base, ier_root):
        menu = []
        for sub in ("MISC", "STRUCTURAL"):
            if _exists(base, ier_root + "\\" + sub):
                menu.append((sub, ier_root + "\\" + sub))
        if menu:
            initial.append({"kind": "dropdown", "label": "IER", "icon": "\U0001F50E", "menu": menu})
        else:
            initial.append({"kind": "button", "label": "IER", "icon": "\U0001F50E",
                            "rel": ier_root, "menu": None})
    schemes = sc("SCHEMES")
    if _exists(base, schemes):
        initial.append({"kind": "button", "label": "SCHEMES", "icon": "\U0001F4D0",
                        "rel": schemes, "menu": None})
    rfi = sc("RFI")
    if _exists(base, rfi):
        menu = [("RFI", rfi), ("-", None)]
        subs = [s for s in ("RFI SENT", "RFI RESPONSE")
                if _exists(base, rfi + "\\" + s)]
        for i, sub in enumerate(subs):
            if _aborted():
                return None
            root = rfi + "\\" + sub
            menu.append((sub, _cascade_or_open(base, root, sub)))
            _report("Reading RFI… {}/{}".format(i + 1, len(subs)))
        initial.append({"kind": "dropdown", "label": "RFI", "icon": "\U0001F4E9", "menu": menu})
    sk = sc("SKETCHES")
    if _exists(base, sk):
        if _aborted():
            return None
        _report("Reading SKETCHES…")
        if _dir_entries(base, sk):
            initial.append({"kind": "dropdown", "label": "SKETCHES", "icon": "\u270F",
                            "menu": _folder_menu(base, sk, label="SKETCHES")})
        else:
            initial.append({"kind": "button", "label": "SKETCHES", "icon": "\u270F",
                            "rel": sk, "menu": None})
    if initial:
        sections.append(("INITIAL", initial))

    # ---------- Submittal stage ----------
    if _aborted():
        return None
    _report("Reading SUBMITTAL…")
    submittal = []
    if _exists(base, sc("SUBMITTAL")):
        submittal.append({"kind": "button", "label": "SUBMITTAL", "icon": "\U0001F4E4",
                          "rel": sc("SUBMITTAL"), "menu": None})
    det_root = sc("DETAILING")
    if _exists(base, det_root):
        menu = [("DETAILING", det_root), ("-", None)]
        subs = [s for s in ("SHOP DRAWINGS", "LAYOUT")
                if _exists(base, det_root + "\\" + s)]
        for i, sub in enumerate(subs):
            if _aborted():
                return None
            root = det_root + "\\" + sub
            # DESIGN COMMENTS (and For Approval / Seq XX levels) live
            # under SHOP DRAWINGS and surface through the recursion.
            menu.append((sub, _cascade_or_open(base, root, sub)))
            _report("Reading DETAILING… {}/{}".format(i + 1, len(subs)))
        submittal.append({"kind": "dropdown", "label": "DETAILING",
                          "icon": "\U0001F3D7", "menu": menu})
    if submittal:
        sections.append(("SUBMITTAL", submittal))

    # ---------- Fabrication stage ----------
    if _aborted():
        return None
    _report("Reading FABRICATION…")
    fab = []
    if _exists(base, sc("APPROVAL SUMMARY")):
        fab.append({"kind": "button", "label": "APPROVAL SUMMARY", "icon": "\U0001F4CA",
                    "rel": sc("APPROVAL SUMMARY"), "menu": None})
    cd_root = sc("CHANGE DOCUMENTS")
    if _exists(base, cd_root):
        cd_target = cd_root + "\\APPROVAL RETURNS"
        if _exists(base, cd_target) and _dir_entries(base, cd_target):
            menu = [("CD", cd_target), ("-", None)]
            names = _dir_entries(base, cd_target)
            for i, name in enumerate(names):
                if _aborted():
                    return None
                rel = cd_target + "\\" + name
                menu.append((name, _cascade_or_open(base, rel, name)))
                _report("Reading CD… {}/{}".format(i + 1, len(names)))
            fab.append({"kind": "dropdown", "label": "CD", "icon": "\U0001F4DD",
                        "menu": menu})
        else:
            fab.append({"kind": "button", "label": "CD", "icon": "\U0001F4DD",
                        "rel": cd_target if _exists(base, cd_target) else cd_root,
                        "menu": None})
    if fab:
        sections.append(("FABRICATION", fab))

    if _aborted():
        return None
    return sections


DATA_DIR = os.path.join(os.environ.get("APPDATA", os.path.expanduser("~")), "IDS", "Navigator")
SETTINGS_FILE = os.path.join(DATA_DIR, "settings.json")
SECTION_CACHE_FILE = os.path.join(DATA_DIR, "section_cache_v3.json")


def load_section_cache():
    """Disk cache of built job sections {job_name: sections}. Lets favorites
    and recently opened jobs load instantly instead of re-scanning the network
    share every time. Invalid entries are simply rebuilt by the caller."""
    try:
        with open(SECTION_CACHE_FILE, "r", encoding="utf-8") as f:
            d = json.load(f)
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def save_section_cache(cache):
    try:
        os.makedirs(DATA_DIR, exist_ok=True)
        tmp = SECTION_CACHE_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(cache, f, indent=2, ensure_ascii=False)
        os.replace(tmp, SECTION_CACHE_FILE)
    except Exception:
        pass

ADMIN_LINKS = [
    {"name": "Connection Ebook", "type": "web",
     "target": "http://192.168.0.118/ids-connection-ebook/dashboard"},
    {"name": "AISC Steel Manual", "type": "file",
     "target": r"\\EgnyteDrive\idsinc\Shared\Engineering\ENGG PHL\REFERENCES\AISC 15th Ed Manual & Specs\AISC Steel Construction Manual, 15th Edition, First Printing Edition, 2017 (HIGH QUALITY VERSION).pdf"},
]

def _admin_links_file():
    # Offline fallback shipped with the app. When embedded in DEA this
    # module lives in client/, so check the app root (parent of client/)
    # first, then the module's own folder (standalone layout).
    here = os.path.dirname(os.path.abspath(__file__))
    if getattr(sys, "frozen", False):
        # Installed exe: next to the exe first, then the PyInstaller bundle
        # (covers running straight from dist\ before install.bat copies the
        # fallback file alongside the exe).
        for base in (os.path.dirname(sys.executable),
                     getattr(sys, "_MEIPASS", None)):
            if base and os.path.isfile(os.path.join(base, "admin_links.json")):
                return os.path.join(base, "admin_links.json")
        return os.path.join(os.path.dirname(sys.executable), "admin_links.json")
    for candidate in (os.path.join(os.path.dirname(here), "admin_links.json"),
                      os.path.join(here, "admin_links.json")):
        if os.path.isfile(candidate):
            return candidate
    return os.path.join(os.path.dirname(here), "admin_links.json")


def _read_links_json(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            clean = [d for d in data
                     if isinstance(d, dict) and d.get("name") and d.get("target")]
            if clean:
                return clean
    except Exception:
        pass
    return None


def _write_links_json(path, links):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(links, f, indent=2, ensure_ascii=False)
    os.replace(tmp, path)


DEA_CONFIG = r"\\EgnyteDrive\idsinc\Shared\Engineering\ENGG PHL\REPORTS\16 Manhour Report LEADERBOARD\DEA App\DEA_Config.xlsx"
SHARED_ADMIN_LINKS = os.path.join(os.path.dirname(DEA_CONFIG), "admin_links.json")


# Built-in icon set for quick-link buttons (Segoe MDL2 Assets glyphs,
# theme-proof unlike color emoji). A link may override its type default
# with any of these via the "Icon..." button in either link manager;
# stored as link["icon"] in settings.json / admin_links.json.
LINK_ICON_CHOICES = [
    ("Globe", "\uE774"),
    ("Document", "\uE8A5"),
    ("Folder", "\uE8B7"),
    ("Calculator", "\uE8EF"),
    ("Mail", "\uE715"),
    ("Link", "\uE71B"),
    ("Star", "\uE734"),
    ("Wrench", "\uE924"),
    ("People", "\uE716"),
    ("Phone", "\uE723"),
    ("Share", "\uE72D"),
    ("Calendar", "\uE787"),
    ("MapPin", "\uE81C"),
    ("Database", "\uE80A"),
    ("Upload", "\uE7E8"),
    ("Download", "\uE896"),
    ("Home", "\uE80F"),
    ("Chart", "\uE825"),
]

# Optional per-link accent color, painted onto the glyph (launch buttons
# get a Pillow-rendered colored image, manager lists get a colored row).
# None = default monochrome look. Stored as link["color"] in
# settings.json / admin_links.json; old links without it are unaffected.
LINK_COLOR_CHOICES = [
    ("Default", None),
    ("Orange", "#F57F17"),
    ("Blue", "#1E88E5"),
    ("Green", "#43A047"),
    ("Red", "#E53935"),
    ("Purple", "#8E24AA"),
    ("Teal", "#00897B"),
]


_LINK_IMAGE_CACHE = {}

def colored_glyph_image(glyph, color, px=18):
    """Render an MDL2 glyph in `color` to a PhotoImage (Pillow, system
    font - exe-safe, no asset file). Cached; Tk images die with their
    last Python reference, hence the module-level cache. Returns None
    if Pillow/the font is unavailable - callers fall back to the plain
    monochrome text glyph."""
    key = (glyph, color, px)
    img = _LINK_IMAGE_CACHE.get(key)
    if img is None:
        try:
            from PIL import Image, ImageDraw, ImageFont, ImageTk
            big = px * 4
            font = ImageFont.truetype("segmdl2.ttf", big)
            im = Image.new("RGBA", (big, big), (0, 0, 0, 0))
            d = ImageDraw.Draw(im)
            bbox = d.textbbox((0, 0), glyph, font=font)
            w, h = bbox[2] - bbox[0], bbox[3] - bbox[1]
            d.text(((big - w) / 2 - bbox[0], (big - h) / 2 - bbox[1]),
                   glyph, font=font, fill=color)
            im = im.resize((px, px), Image.LANCZOS)
            img = ImageTk.PhotoImage(im)
            _LINK_IMAGE_CACHE[key] = img
        except Exception:
            return None
    return img


def default_icon_for_type(typ):
    return {"web": "\uE774", "file": "\uE8A5", "folder": "\uE8B7"}.get(typ, "\uE71B")


_ORANGE_DOT_CACHE = {}

def orange_dot_image(size=14, color="#F57F17"):
    """Small orange-circle PhotoImage for buttons. Drawn at runtime (no
    asset file, so exe-safe) and cached - Tk images die with their last
    Python reference, hence the module-level cache."""
    key = (size, color)
    img = _ORANGE_DOT_CACHE.get(key)
    if img is None:
        try:
            img = tk.PhotoImage(width=size, height=size)
            r = size / 2 - 1
            for y in range(size):
                for x in range(size):
                    dx, dy = x + 0.5 - size / 2, y + 0.5 - size / 2
                    if dx * dx + dy * dy <= r * r:
                        img.put(color, (x, y))
                    else:
                        img.transparency_set(x, y, True)
            _ORANGE_DOT_CACHE[key] = img
        except Exception:
            return None
    return img


def resolve_link_icon(link):
    """Explicit per-link icon if one was picked, else the type default."""
    return link.get("icon") or default_icon_for_type(link.get("type"))


def pick_link_icon(parent, current_glyph=None, current_color=None):
    """Modal glyph + color picker. Returns (glyph, color) on Save (color
    None = default monochrome), or None if cancelled. Opens preselected
    on the link's current choices, so saving untouched keeps them."""
    result = {"glyph": current_glyph, "color": current_color, "saved": False}
    win = tk.Toplevel(parent)
    win.title("Choose Icon")
    win.resizable(False, False)
    win.transient(parent)
    win.grab_set()
    summary = tk.StringVar()
    dots = []

    def refresh_summary():
        gname = next((n for n, g in LINK_ICON_CHOICES if g == result["glyph"]),
                     "Type default" if not result["glyph"] else "Custom")
        cname = next((n for n, c in LINK_COLOR_CHOICES if c == result["color"]),
                     "Custom" if result["color"] else "Default")
        summary.set("Icon: {}    Color: {}".format(gname, cname))

    ttk.Label(win, text="Choose an icon:", font=("Segoe UI", 10, "bold")).pack(
        padx=14, pady=(12, 6))
    grid = ttk.Frame(win)
    grid.pack(padx=14, pady=(0, 6))
    for i, (name, glyph) in enumerate(LINK_ICON_CHOICES):
        ttk.Button(grid, text="{}  {}".format(glyph, name), width=16,
                   command=lambda g=glyph: (result.update(glyph=g), refresh_summary())
                   ).grid(row=i // 3, column=i % 3, padx=4, pady=4)
    ttk.Label(win, text="Color:", font=("Segoe UI", 10, "bold")).pack(padx=14, pady=(6, 2))
    crow = ttk.Frame(win)
    crow.pack(padx=14, pady=(0, 6))
    for cname, chex in LINK_COLOR_CHOICES:
        if chex is None:
            ttk.Button(crow, text=cname,
                       command=lambda: (result.update(color=None), refresh_summary())
                       ).pack(side="left", padx=4)
        else:
            dot = orange_dot_image(12, chex)
            dots.append(dot)
            ttk.Button(crow, text=cname, image=dot, compound="left",
                       command=lambda c=chex: (result.update(color=c), refresh_summary())
                       ).pack(side="left", padx=4)
    ttk.Label(win, textvariable=summary, font=("Segoe UI", 9),
              foreground="#68727d").pack(padx=14, pady=(0, 6))
    brow = ttk.Frame(win)
    brow.pack(pady=(0, 12))
    ttk.Button(brow, text="Save",
               command=lambda: (result.update(saved=True), win.destroy())).pack(side="left")
    ttk.Button(brow, text="Cancel", command=win.destroy).pack(side="left", padx=(6, 0))
    refresh_summary()
    win.bind("<Escape>", lambda e: win.destroy())
    parent.wait_window(win)
    if result["saved"]:
        return result["glyph"], result["color"]
    return None


_ADMIN_LINKS_CACHE = {"links": None, "at": 0.0}
_ADMIN_LINKS_TTL_SECONDS = 60.0


def invalidate_admin_links_cache():
    """Forget the memoized admin links so the next read hits disk."""
    _ADMIN_LINKS_CACHE["links"] = None
    _ADMIN_LINKS_CACHE["at"] = 0.0


def get_admin_links():
    """Admin-defined quick links, shared centrally so every install sees the
    same set. Priority:
      1. admin_links.json sitting next to the shared DEA_Config.xlsx
      2. a local admin_links.json shipped with the app (offline fallback)
      3. the embedded defaults above.

    Memoized briefly: the click path calls this several times per render,
    and each uncached call hits the network share."""
    now = time.monotonic()
    if (_ADMIN_LINKS_CACHE["links"] is not None
            and now - _ADMIN_LINKS_CACHE["at"] < _ADMIN_LINKS_TTL_SECONDS):
        return _ADMIN_LINKS_CACHE["links"]
    links = _read_links_json(SHARED_ADMIN_LINKS)
    if not links:
        links = _read_links_json(_admin_links_file())
    if not links:
        links = ADMIN_LINKS
    _ADMIN_LINKS_CACHE["links"] = links
    _ADMIN_LINKS_CACHE["at"] = now
    return links


def load_admin_credentials():
    """Username/password that unlock the admin quick-link editor, read live
    from the 'App Settings' sheet of the shared DEA_Config.xlsx. Goes
    through the shared config opener so a file-open-encrypted config still
    works (falls back to a plain openpyxl read when run standalone)."""
    try:
        try:
            from shared.config_manager import open_config_workbook
            wb = open_config_workbook(DEA_CONFIG)
        except ImportError:
            import openpyxl
            wb = openpyxl.load_workbook(DEA_CONFIG, data_only=True, read_only=True)
        try:
            if "App Settings" not in wb.sheetnames:
                return None
            creds = {}
            for row in wb["App Settings"].iter_rows(values_only=True):
                if row and row[0] in ("AdminUsername", "AdminPassword") and row[1] not in (None, ""):
                    creds[row[0]] = str(row[1]).strip()
            if creds.get("AdminUsername") and creds.get("AdminPassword"):
                return creds["AdminUsername"], creds["AdminPassword"]
        finally:
            wb.close()
    except Exception:
        pass
    return None


def _user_links(links):
    """Per-user quick links, with any admin-defined links removed so they
    never appear twice (admin links are always shown separately)."""
    admin_targets = {l.get("target") for l in get_admin_links()}
    return [l for l in (links or []) if l.get("target") not in admin_targets]

def ensure_data():
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(SETTINGS_FILE):
        save_settings({"favorites": [], "recent": [], "links": []})

def load_settings():
    ensure_data()
    try:
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        data["links"] = _user_links(data.get("links"))
        return data
    except Exception:
        return {"favorites": [], "recent": [], "links": []}

def save_settings(data):
    os.makedirs(DATA_DIR, exist_ok=True)
    tmp = SETTINGS_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    os.replace(tmp, SETTINGS_FILE)

def normalize_path(path):
    """Tidy a path before handing it to the OS: strip stray quotes/whitespace
    (including newlines from pasted Explorer address bars), unify slashes,
    and collapse doubled backslashes -- keeps UNC quick links working."""
    if not path:
        return path
    path = path.strip().strip('"').strip("'")
    is_unc = path.startswith("\\\\") or path.startswith("//")
    path = path.replace("/", "\\")
    if is_unc:
        rest = path.lstrip("\\")
        rest = re.sub(r"\\{2,}", r"\\", rest)
        path = "\\\\" + rest
    else:
        path = re.sub(r"\\{2,}", r"\\", path)
    if len(path) > 3 and path.endswith("\\"):
        path = path.rstrip("\\")
    return path


def open_target(target):
    t = normalize_path(target)
    if not t:
        return
    try:
        os.startfile(t)
    except Exception as e:
        messagebox.showerror(APP_NAME, "Could not open:\n{}\n\n{}".format(t, e))


FILE_EXTS = (".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
             ".txt", ".dwg", ".zip", ".exe", ".msg", ".eml")

def detect_link_type(target):
    """Figure out the link type from its target instead of asking the user."""
    raw = (target or "").strip()
    if raw.lower().startswith(("http://", "https://")):
        return "web"
    t = normalize_path(target)
    if os.path.splitext(t)[1].lower() in FILE_EXTS:
        return "file"
    try:
        if os.path.isfile(t):
            return "file"
    except OSError:
        pass
    return "folder"


def parse_drop_paths(data):
    """Split Explorer's drag-drop payload into raw paths. Explorer wraps paths
    containing spaces in {braces}; a drop may carry several items."""
    paths = []
    i, n = 0, len(data)
    while i < n:
        if data[i] == "{":
            j = data.find("}", i)
            if j == -1:
                break
            paths.append(data[i + 1:j])
            i = j + 1
        elif data[i] in " \t":
            i += 1
        else:
            j = i
            while j < n and data[j] not in " \t":
                j += 1
            paths.append(data[i:j])
            i = j
    return [p for p in paths if p]


def to_unc(path):
    """Resolve a dropped path to its UNC form: drive letters (mapped drives)
    become \\\\server\\share via PowerShell, UNC paths pass through."""
    p = (path or "").strip().strip("{}").strip().replace("/", "\\")
    if not p:
        return ""
    if p.startswith("\\\\"):
        return normalize_path(p)
    m = re.match(r"^([A-Za-z]):(.*)$", p)
    if not m:
        return p
    drive, rest = m.group(1).upper(), m.group(2)
    try:
        out = subprocess.check_output(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command",
             "(Get-PSDrive {}).DisplayRoot".format(drive)],
            stderr=subprocess.DEVNULL, text=True).strip()
    except Exception:
        out = ""
    if out:
        return out.rstrip("\\") + rest
    return p

class Navigator(tk.Frame):
    def __init__(self, master=None, **kwargs):
        super().__init__(master, **kwargs)
        self.settings = load_settings()
        self.jobs = []
        self.current_job = None
        self.filtered_jobs = []
        self._folder_cache = load_section_cache()
        self._cache_lock = threading.Lock()
        self._collapsed = {"MANAGE LINKS"}
        self._hover_menus = []
        self._job_token = 0  # guards async renders against rapid re-clicks
        self._select_after_id = None  # debounce for arrow-key scrolling
        self._build_style()
        self._build_ui()
        self.refresh_jobs()
        self.refresh_favorites()
        self._prewarm()

    def _build_style(self):
        # NOTE: ttk styles (and especially theme_use) are application-global.
        # This tab must NEVER call theme_use() or restyle stock selectors
        # ("Treeview", "TButton", ...) - doing so repainted the whole DEA
        # window in clam's gray/brown. Every selector below is Nav-prefixed
        # so only widgets inside this tab are affected; DEA keeps its
        # initial native look.
        style = ttk.Style(self)
        style.configure("App.TFrame", background="#f5f6f8")
        style.configure("Sidebar.TFrame", background="#20252b")
        style.configure("Sidebar.TLabel", background="#20252b", foreground="#e8edf2")
        style.configure("Title.TLabel", font=("Segoe UI", 15, "bold"), background="#f5f6f8", foreground="#20252b")
        style.configure("Sub.TLabel", font=("Segoe UI", 9), background="#f5f6f8", foreground="#68727d")

        border = dict(borderwidth=1, relief="solid", bordercolor="#c2c8cf",
                      lightcolor="#c2c8cf", darkcolor="#c2c8cf")
        style.configure("Card.TButton", font=("Segoe UI", 10), padding=(12, 9),
                        background="#ffffff", foreground="#20252b", **border)
        style.map("Card.TButton",
                  background=[("active", "#eef2f6"), ("pressed", "#dfe5ea")],
                  bordercolor=[("active", "#b0b8c2")])
        style.configure("Card.TMenubutton", font=("Segoe UI", 10), padding=(12, 9),
                        background="#ffffff", foreground="#20252b", **border)
        style.map("Card.TMenubutton",
                  background=[("active", "#eef2f6"), ("pressed", "#dfe5ea")],
                  bordercolor=[("active", "#b0b8c2")])
        style.configure("Link.TButton", font=("Segoe UI", 10), padding=(12, 9))
        style.configure("Toggle.TButton", font=("Segoe UI", 9), padding=(4, 2),
                        borderwidth=0, relief="flat", background="#f5f6f8")
        style.map("Toggle.TButton", background=[("active", "#e6e8ec")])
        style.configure("Nav.Treeview", rowheight=28, font=("Segoe UI", 9),
                        background="#ffffff", fieldbackground="#ffffff")
        style.configure("Nav.Treeview.Heading", font=("Segoe UI", 9, "bold"))

    def _build_ui(self):
        self.configure(bg="#f5f6f8")
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        sidebar = ttk.Frame(self, style="Sidebar.TFrame", padding=(14, 16))
        sidebar.grid(row=0, column=0, sticky="ns")
        sidebar.configure(width=265)
        sidebar.grid_propagate(False)
        self._sidebar = sidebar

        ttk.Label(sidebar, text="NAVITOOL 2.0", style="Sidebar.TLabel",
                  font=("Segoe UI", 14, "bold")).pack(anchor="w", pady=(0, 18))

        ttk.Label(sidebar, text="FAVORITES", style="Sidebar.TLabel",
                  font=("Segoe UI", 8, "bold")).pack(anchor="w")
        fav_frame = tk.Frame(sidebar, bg="#20252b")
        fav_frame.pack(fill="x", pady=(6, 10))
        self.fav_list = tk.Listbox(fav_frame, bg="#20252b", fg="#e8edf2",
                                   selectbackground="#3d79b9", selectforeground="white",
                                   relief="flat", highlightthickness=0,
                                   font=("Segoe UI", 9), activestyle="none",
                                   height=10, exportselection=False)
        fav_scroll = tk.Scrollbar(fav_frame, orient="vertical", command=self.fav_list.yview,
                                  bg="#2b3138", troughcolor="#20252b",
                                  activebackground="#3d79b9", relief="flat",
                                  borderwidth=0, width=12, highlightthickness=0)
        self.fav_list.configure(yscrollcommand=fav_scroll.set)
        self.fav_list.pack(side="left", fill="x", expand=True)
        fav_scroll.pack(side="right", fill="y")
        self.fav_list.bind("<<ListboxSelect>>", self._on_favorite_select)
        self.fav_list.bind("<Double-Button-1>", lambda e: self._favorite_open_folder())
        self.fav_list.bind("<Button-3>", self._favorite_context_menu)

        ttk.Label(sidebar, text="JOBS", style="Sidebar.TLabel",
                  font=("Segoe UI", 8, "bold")).pack(anchor="w")
        self.job_search = ttk.Entry(sidebar)
        self.job_search.pack(fill="x", pady=(6, 10))
        self.job_search.insert(0, "Search jobs...")
        self.job_search.configure(foreground="#7a828a")
        self.job_search.bind("<FocusIn>", self._search_focus_in)
        self.job_search.bind("<FocusOut>", self._search_focus_out)
        self.job_search.bind("<KeyRelease>", lambda e: self.filter_jobs())

        self.job_list = tk.Listbox(sidebar, bg="#20252b", fg="#e8edf2",
                                   selectbackground="#3d79b9", selectforeground="white",
                                   relief="flat", highlightthickness=0,
                                   font=("Segoe UI", 9), activestyle="none",
                                   exportselection=False)
        self.job_list.pack(fill="both", expand=True)
        self.job_list.bind("<<ListboxSelect>>", self.select_job)
        self.job_list.bind("<Double-Button-1>", lambda e: self.open_current_job())
        self.job_list.bind("<Button-3>", self._job_context_menu)

        ttk.Button(sidebar, text="\u21bb Refresh Jobs", command=self.refresh_jobs).pack(fill="x", pady=(10, 5))

        main = ttk.Frame(self, style="App.TFrame", padding=(22, 18))
        main.grid(row=0, column=1, sticky="nsew")
        main.grid_columnconfigure(0, weight=1)
        main.grid_rowconfigure(4, weight=1)

        top = ttk.Frame(main, style="App.TFrame")
        top.grid(row=0, column=0, sticky="ew")
        top.grid_columnconfigure(0, weight=1)
        self.title_var = tk.StringVar(value="Select a job")
        ttk.Label(top, textvariable=self.title_var, style="Title.TLabel").grid(row=0, column=0, sticky="w")
        self.path_var = tk.StringVar(value=JOB_ROOT)
        ttk.Label(top, textvariable=self.path_var, style="Sub.TLabel").grid(row=1, column=0, sticky="w", pady=(3, 0))

        actions = ttk.Frame(main, style="App.TFrame")
        actions.grid(row=1, column=0, sticky="ew", pady=(10, 4))
        ttk.Button(actions, text="Open Job", command=self.open_current_job).pack(side="left")
        ttk.Button(actions, text="Add to Favorites", command=self.add_favorite).pack(side="left", padx=6)
        self._quick_dot = orange_dot_image()
        ttk.Button(actions, text=" Quick Links", image=self._quick_dot,
                   compound="left", command=self._jump_to_quick).pack(side="left")

        self.status_var = tk.StringVar(value="")
        status = ttk.Label(main, textvariable=self.status_var, style="Sub.TLabel")
        status.grid(row=3, column=0, sticky="w", pady=(0, 4))

        def _sync_status_row(*_args):
            # The status line only takes up space while it has something to
            # say - an empty message collapses the row entirely.
            if self.status_var.get().strip():
                status.grid()
            else:
                status.grid_remove()

        self.status_var.trace_add("write", _sync_status_row)
        _sync_status_row()

        # Job-folders loading row: shown only while a background section
        # build runs (cache misses) - phase text, indeterminate bar (a
        # determinate total would cost a full pre-walk just for counting),
        # and a Cancel button. Mirrors the Latest Details progress row.
        prog = ttk.Frame(main, style="App.TFrame")
        prog.grid(row=2, column=0, sticky="ew", pady=(0, 4))
        self._folders_prog_label = ttk.Label(prog, text="", style="Sub.TLabel",
                                             font=("Segoe UI", 8))
        self._folders_prog_label.pack(side="left")
        self._folders_cancel_btn = ttk.Button(prog, text="Cancel",
                                              command=self._cancel_folders_build)
        self._folders_prog_bar = ttk.Progressbar(prog, mode="indeterminate",
                                                 length=180)
        self._folders_prog_bar.pack(side="right", padx=(6, 0))
        self._folders_prog = prog
        self._folders_cancel = None
        self._set_folders_progress_visible(False)

        # Fixed tab strip + one scroll region per page: the Job Folders /
        # Latest Details / Quick Links tabs never scroll away; each page
        # scrolls its own content (exactly one canvas per scrolling page,
        # so the mouse wheel never fights nested scrollers). The Latest
        # page holds only the panel, which freezes its own header and
        # family tabs above its rows canvas (see _mount_latest_panel).
        content_wrap = ttk.Frame(main, style="App.TFrame")
        content_wrap.grid(row=4, column=0, sticky="nsew")
        content_wrap.grid_columnconfigure(0, weight=1)
        content_wrap.grid_rowconfigure(1, weight=1)
        self._sub_nb = ttk.Notebook(content_wrap)
        self._sub_nb.grid(row=0, column=0, sticky="ew")
        self._page_frames = {}
        self._page_canvas = {}
        self._page_inner = {}
        self._scroll_map = {}
        for _pname, _plabel in (("folders", "Job Folders"),
                                ("latest", "Latest Details"),
                                ("quick", "Quick Links")):
            _pg = ttk.Frame(self._sub_nb, style="App.TFrame")
            self._sub_nb.add(_pg, text=_plabel)
            self._page_frames[_pname] = _pg
            if _pname == "latest":
                self._page_latest = _pg
                continue
            if _pname == "quick":
                self._page_quick = _pg
            _pg.grid_columnconfigure(0, weight=1)
            _pg.grid_rowconfigure(0, weight=1)
            _cv = tk.Canvas(_pg, background="#f5f6f8", highlightthickness=0)
            _vsb = ttk.Scrollbar(_pg, orient="vertical", command=_cv.yview)
            _inner = ttk.Frame(_cv, style="App.TFrame")
            _cv.create_window((0, 0), window=_inner, anchor="nw")
            _cv.configure(yscrollcommand=_vsb.set)
            _cv.grid(row=0, column=0, sticky="nsew")
            _vsb.grid(row=0, column=1, sticky="ns")
            for _col in range(3):
                _inner.grid_columnconfigure(_col, weight=1, uniform="btn")
            _inner.bind("<Configure>",
                        lambda e, c=_cv: c.configure(scrollregion=c.bbox("all")))
            _cv.bind("<Configure>", lambda e, c=_cv: self._fit_canvas_width(c))
            self._page_canvas[_pname] = _cv
            self._page_inner[_pname] = _inner
            self._scroll_map[_inner] = _cv
        self._latest_panel = None
        self._links_tree = None
        self._links_notice = ""
        self._link_images = []
        self.bind_all("<MouseWheel>", self._scroll_content_wheel, add="+")

        self._show_placeholder()
        self._fit_sidebar_width()

    def _fit_sidebar_width(self):
        """Size the sidebar to its content, in real pixels on this machine
        (a fixed width clips names under display scaling). Widest favorite
        / job text plus padding, clamped so one long name can't eat the
        window. Re-run when the job list changes, not on every keystroke
        of the search box (stability over twitchiness)."""
        try:
            import tkinter.font as tkfont
            font = tkfont.Font(root=self, family="Segoe UI", size=9)
            texts = list(self.job_list.get(0, "end")) + \
                list(self.fav_list.get(0, "end"))
            widest = max([font.measure(t) for t in texts if t] + [0])
            width = min(max(widest + 28 + 14, 220), 360)
            self._sidebar.configure(width=width)
        except Exception:
            pass

    def _fit_canvas_width(self, canvas):
        """Keep a page canvas's inner frame as wide as the canvas so the
        3-column button grid always fills the visible width."""
        try:
            width = canvas.winfo_width()
            for item in canvas.find_all():
                canvas.itemconfigure(item, width=width)
        except tk.TclError:
            pass

    def _search_focus_in(self, _):
        if self.job_search.get() == "Search jobs...":
            self.job_search.delete(0, "end")
            self.job_search.configure(foreground="#20252b")

    def _search_focus_out(self, _):
        if not self.job_search.get():
            self.job_search.insert(0, "Search jobs...")
            self.job_search.configure(foreground="#7a828a")

    def refresh_jobs(self):
        # NOTE: the section cache is deliberately NOT wiped here. A full
        # wipe made every post-refresh click re-scan the network share
        # (the main click-lag complaint). Only entries for jobs that no
        # longer exist are pruned; favorites/recents rebuild in the
        # background via _prewarm(force=True); and the currently open job
        # (if any) is rebuilt immediately below - so a rename or a new
        # folder shows up as soon as you press this button.
        try:
            names = []
            with os.scandir(JOB_ROOT) as it:
                for entry in it:
                    if entry.is_dir():
                        names.append(entry.name)
            self.jobs = sorted(names, key=str.lower)
        except Exception as e:
            self.jobs = []
            self._show_placeholder("Cannot access JOBS directory.\n\n{}\n\n{}".format(JOB_ROOT, e))
            return
        alive = set(self.jobs)
        with self._cache_lock:
            stale = [k for k in self._folder_cache if k not in alive]
            for k in stale:
                del self._folder_cache[k]
            if stale:
                save_section_cache(self._folder_cache)
        self.filter_jobs()
        self._fit_sidebar_width()
        self._prewarm(force=True)
        if self.current_job and self.current_job in self.jobs:
            with self._cache_lock:
                self._folder_cache.pop(self.current_job, None)
            self._show_job_async()
            try:
                import datetime as _dt
                self.status_var.set("Refreshed %s." % _dt.datetime.now().strftime("%H:%M"))
            except Exception:
                pass

    def filter_jobs(self):
        q = self.job_search.get().strip().lower()
        if q == "search jobs...":
            q = ""
        self.filtered_jobs = [j for j in self.jobs if q in j.lower()]
        self.job_list.delete(0, "end")
        for job in self.filtered_jobs:
            self.job_list.insert("end", job)

    def select_job(self, _=None):
        # Debounced: arrow-key scrolling fires <<ListboxSelect>> on every
        # step, and each step used to trigger a full synchronous network
        # scan. Wait briefly so only the settled-on job actually renders.
        if self._select_after_id is not None:
            try:
                self.after_cancel(self._select_after_id)
            except tk.TclError:
                pass
            self._select_after_id = None

        def fire():
            self._select_after_id = None
            sel = self.job_list.curselection()
            if not sel:
                return
            self._open_job(self.filtered_jobs[sel[0]])

        self._select_after_id = self.after(80, fire)

    def _open_job(self, name):
        self.current_job = name
        self.title_var.set(name)
        self.path_var.set(os.path.join(JOB_ROOT, name))
        self._show_job_async()
        self._add_recent(name)
        self.status_var.set("")
        if name in self.filtered_jobs:
            idx = self.filtered_jobs.index(name)
            self.job_list.selection_clear(0, "end")
            self.job_list.selection_set(idx)
            self.job_list.see(idx)

    def _clear_content(self):
        for inner in self._page_inner.values():
            for w in inner.winfo_children():
                w.destroy()
        for w in self._page_frames["latest"].winfo_children():
            w.destroy()
        self._latest_panel = None
        for canvas in list(self._page_canvas.values()):
            try:
                canvas.yview_moveto(0.0)
            except (tk.TclError, AttributeError):
                pass

    def _scroll_content_wheel(self, event):
        """Mouse-wheel scrolling for the page canvases. Bound app-wide
        (bind_all) but only acts when the pointer is actually over one of
        this tab's scroll regions (a page inner or the Latest panel's
        rows), so it never hijacks scrolling in DEA's other tabs
        or dialogs."""
        try:
            widget = self.winfo_containing(event.x_root, event.y_root)
        except tk.TclError:
            return
        node = widget
        try:
            target = None
            while node is not None:
                if node in self._scroll_map:
                    target = self._scroll_map[node]
                    break
                node = node.master
        except tk.TclError:
            return
        if target is None:
            return
        try:
            target.yview_scroll(-1 * (event.delta // 120), "units")
        except tk.TclError:
            pass

    def _show_placeholder(self, text="Select a job from the left."):
        self._clear_content()
        page_folders = self._page_inner["folders"]
        page_latest = self._page_frames["latest"]
        page_quick = self._page_inner["quick"]
        ttk.Label(page_folders, text=text, style="Sub.TLabel",
                  font=("Segoe UI", 11)).grid(row=0, column=0, columnspan=3, pady=(30, 10))
        # No job selected: the folders page is just the placeholder, quick
        # links live on their own tab, and the Latest page explains itself
        # (the panel guides job selection).
        all_links = get_admin_links() + self.settings.get("links", [])
        width = max([len(l["name"]) + 3 for l in all_links] + [1]) + 2
        self._render_quick_page(width)
        self._mount_latest_panel(page_latest)

    # ---------- main content ----------

    def _get_sections(self, name):
        """Return built job sections, using the in-memory/disk cache when
        available and persisting newly built ones to disk."""
        with self._cache_lock:
            if name in self._folder_cache:
                return self._folder_cache[name]
        try:
            sections = build_job_sections(name)
        except Exception:
            return []
        with self._cache_lock:
            self._folder_cache[name] = sections
            save_section_cache(self._folder_cache)
        return sections

    def _prewarm(self, force=False):
        """Pre-build sections for favorites AND recently opened jobs in a
        background thread so they are already cached when the user clicks
        them. With force=True (Refresh Jobs) existing cache is rebuilt so
        data stays current."""
        targets = []
        seen = set()
        for n in list(self.settings.get("favorites", [])) + list(self.settings.get("recent", [])):
            if n and n not in seen:
                seen.add(n)
                with self._cache_lock:
                    if force or n not in self._folder_cache:
                        targets.append(n)
        if not targets:
            return

        def work():
            for name in targets:
                base = os.path.join(JOB_ROOT, name)
                if not os.path.isdir(base):
                    continue
                try:
                    sections = build_job_sections(name)
                except Exception:
                    continue
                with self._cache_lock:
                    self._folder_cache[name] = sections
                    save_section_cache(self._folder_cache)

        threading.Thread(target=work, daemon=True).start()

    def _toggle_section(self, key, header, frame, label):
        """Collapse/expand a section. `key` tracks the state in _collapsed;
        `label` is the text shown next to the arrow ("" for the untitled
        stage groups, "QUICK LINKS" for the links block)."""
        suffix = "  " + label if label else ""
        if frame.winfo_ismapped():
            frame.grid_remove()
            header.configure(text="\u25B8" + suffix)
            self._collapsed.add(key)
        else:
            frame.grid()
            header.configure(text="\u25BE" + suffix)
            self._collapsed.discard(key)

    def _show_job(self):
        """Synchronous render from cache (or a fast local build). Used by
        the quick-link editors, which only need to repaint the links."""
        self._clear_content()
        if not self.current_job:
            self._show_placeholder()
            return
        self._render_sections(self._get_sections(self.current_job))

    def _set_folders_progress_visible(self, visible):
        try:
            if visible:
                self._folders_prog.grid()
                self._folders_cancel_btn.pack(side="right")
                self._folders_prog_bar.pack(side="right", padx=(6, 0))
                try:
                    self._folders_prog_bar.start(50)
                except tk.TclError:
                    pass
                self._folders_prog_label.config(text="Starting…")
            else:
                try:
                    self._folders_prog_bar.stop()
                except tk.TclError:
                    pass
                self._folders_prog.grid_remove()
        except tk.TclError:
            pass

    def _update_folders_progress(self, token, text):
        if token != self._job_token:
            return
        try:
            self._folders_prog_label.config(text=text)
        except tk.TclError:
            pass

    def _cancel_folders_build(self):
        if self._folders_cancel is not None:
            try:
                self._folders_cancel.set()
            except Exception:
                pass

    def _show_job_async(self):
        """Click path: paint instantly, never block the UI on the network.

        - Cache hit  -> render immediately (no thread, no flicker).
        - Cache miss -> show a "Loading..." placeholder at once, build the
          sections in a background thread with a progress row (cancellable),
          then paint via after(). A token guard drops stale results when
          the user clicks another job before the build finishes."""
        self._clear_content()
        if not self.current_job:
            self._show_placeholder()
            return
        with self._cache_lock:
            cached = self._folder_cache.get(self.current_job)
        if cached is not None:
            self._render_sections(cached)
            return
        self._job_token += 1
        token = self._job_token
        name = self.current_job
        if self._folders_cancel is not None:
            try:
                self._folders_cancel.set()
            except Exception:
                pass
        cancel = threading.Event()
        self._folders_cancel = cancel
        self._clear_content()
        ttk.Label(self._page_inner["folders"], text="Loading '{}'...".format(name),
                  style="Sub.TLabel", font=("Segoe UI", 11)).grid(
                      row=0, column=0, columnspan=3, pady=(30, 10))
        self._set_folders_progress_visible(True)

        def on_progress(text):
            try:
                self.after(0, lambda: self._update_folders_progress(token, text))
            except Exception:
                pass

        def work():
            try:
                sections = build_job_sections(name, progress=on_progress,
                                              cancel=cancel.is_set)
            except Exception:
                sections = []
            try:
                self.after(0, lambda: self._finish_folders(token, name, sections))
            except Exception:
                pass

        threading.Thread(target=work, daemon=True).start()

    def _finish_folders(self, token, name, sections):
        if token != self._job_token or self.current_job != name:
            return  # user moved on; drop this stale result
        self._set_folders_progress_visible(False)
        if sections is None:
            # Cancelled - keep the placeholder, note how to retry.
            self.status_var.set("Loading cancelled - click the job again to retry.")
            return
        try:
            if not self.winfo_exists():
                return
        except tk.TclError:
            return
        with self._cache_lock:
            self._folder_cache[name] = sections
            save_section_cache(self._folder_cache)
        self._clear_content()
        self._render_sections(sections)

    def _mount_latest_panel(self, page):
        """Latest Details panel filling the Latest page (fixed header +
        family tabs above its own rows canvas - the page itself does not
        scroll). Lazy import avoids a circular import (that module imports
        JOB_ROOT etc. from here). Recreated per render; instant cache
        paint keeps it free, and its token guard drops stale background
        results."""
        from client.latest_details_tab import LatestDetailsTab
        self._latest_panel = LatestDetailsTab(page, lambda: self.current_job,
                                              embedded=True)
        self._latest_panel.pack(fill="both", expand=True)
        rows_canvas = getattr(self._latest_panel, "_rows_canvas", None)
        self._scroll_map = {k: v for k, v in self._scroll_map.items()
                            if k in self._page_inner.values()}
        if rows_canvas is not None:
            try:
                rows_inner = self._latest_panel.body
                self._scroll_map[rows_inner] = rows_canvas
            except AttributeError:
                pass
        self._latest_panel.refresh()

    def _render_sections(self, sections):

        width = 1
        for title, items in sections:
            for it in items:
                w = len(it["label"]) + (2 if it["kind"] == "dropdown" else 3)
                width = max(width, w)
        all_links = get_admin_links() + self.settings.get("links", [])
        for link in all_links:
            width = max(width, len(link["name"]) + 3)
        width += 2

        row = 0
        page_folders = self._page_inner["folders"]
        page_latest = self._page_frames["latest"]
        # Stage groups keep their collapsible arrow-toggle, but the toggle
        # shows ONLY the arrow - no title text. Only QUICK LINKS keeps a
        # titled header (see _render_quick_links). Each group still gets
        # its own row-block so its buttons start on a fresh row.
        for title, items in sections:
            # Plain label, not a button: a button outline can never appear,
            # and the hand cursor still says "click me". Still collapses.
            toggle = tk.Label(page_folders, text="\u25BE", font=("Segoe UI", 9),
                              bg="#f5f6f8", fg="#555555", cursor="hand2")
            toggle.grid(row=row, column=0, columnspan=3, sticky="w", padx=8,
                        pady=(2 if row == 0 else 10, 0))
            row += 1
            frame = ttk.Frame(page_folders, style="App.TFrame")
            frame.grid(row=row, column=0, columnspan=3, sticky="ew", padx=2)
            self._render_items(frame, items, width)
            toggle.bind("<Button-1>", lambda e, t=title, h=toggle, f=frame:
                        self._toggle_section(t, h, f, ""))
            if title in self._collapsed:
                frame.grid_remove()
                toggle.configure(text="\u25B8")
            row += 1

        all_links = get_admin_links() + self.settings.get("links", [])
        self._render_quick_page(width)
        self._mount_latest_panel(page_latest)

    def _jump_to_quick(self):
        """'Quick Links' beside Add to Favorites: flip to the Quick Links
        sub-tab - launch buttons on top, the link manager below."""
        try:
            self._sub_nb.select(self._page_quick)
        except (tk.TclError, AttributeError):
            pass

    def _render_quick_page(self, width):
        """Whole Quick Links tab: launch buttons plus the embedded link
        manager underneath."""
        self._links_tree = None
        self._link_images = []
        page_quick = self._page_inner["quick"]
        self._render_quick_links(page_quick, width, 0)
        self._render_links_manager(page_quick, 2)

    def _refresh_quick_links(self, select=None):
        """Repaint only the Quick Links tab after a manager edit - the Job
        Folders and Latest pages keep their state, and `select` (a tree
        iid) restores the edited row's selection in the rebuilt list."""
        page_quick = self._page_inner["quick"]
        for w in page_quick.winfo_children():
            w.destroy()
        all_links = get_admin_links() + self.settings.get("links", [])
        width = max([len(l["name"]) + 3 for l in all_links] + [1]) + 2
        self._render_quick_page(width)
        if select is not None and self._links_tree is not None:
            try:
                iid = str(select)
                if iid in self._links_tree.get_children():
                    self._links_tree.selection_set(iid)
                    self._links_tree.see(iid)
            except tk.TclError:
                pass

    def _render_quick_links(self, host, width, row):
        """Quick-links card block with its own collapse toggle. Shown under a
        selected job and on the empty (no-job-selected) screen."""
        all_links = get_admin_links() + self.settings.get("links", [])
        links_toggle = tk.Label(host, text="\u25BE  QUICK LINKS",
                                font=("Segoe UI", 9, "bold"),
                                bg="#f5f6f8", fg="#20252b", cursor="hand2")
        links_toggle.grid(row=row, column=0, columnspan=3, sticky="w", padx=8, pady=(14, 0))
        row += 1
        links_frame = ttk.Frame(host, style="App.TFrame")
        links_frame.grid(row=row, column=0, columnspan=3, sticky="ew", padx=2)
        for j, link in enumerate(all_links):
            color = link.get("color")
            glyph_img = colored_glyph_image(resolve_link_icon(link), color) if color else None
            if glyph_img is not None:
                self._link_images.append(glyph_img)
                ttk.Button(links_frame, text=link["name"], image=glyph_img,
                           compound="left", style="Card.TButton", width=width,
                           command=lambda x=link["target"]: open_target(x)).grid(
                               row=j // 3, column=j % 3, sticky="ew", padx=5, pady=5)
            else:
                ttk.Button(links_frame, text="{}  {}".format(
                    resolve_link_icon(link), link["name"]),
                    style="Card.TButton", width=width,
                    command=lambda x=link["target"]: open_target(x)).grid(
                        row=j // 3, column=j % 3, sticky="ew", padx=5, pady=5)
        links_frame.grid_columnconfigure(0, weight=1, uniform="btn")
        links_frame.grid_columnconfigure(1, weight=1, uniform="btn")
        links_frame.grid_columnconfigure(2, weight=1, uniform="btn")
        links_toggle.bind("<Button-1>", lambda e, h=links_toggle, f=links_frame:
                            self._toggle_section("QUICK LINKS", h, f, "QUICK LINKS"))
        if "QUICK LINKS" in self._collapsed:
            links_frame.grid_remove()
            links_toggle.configure(text="\u25B8  QUICK LINKS")
        return row + 1

    def _fill_menu(self, menu, entries):
        """Populate a tk.Menu from nested [(label, rel-or-children), ...]
        entries to any depth: "-" is a separator, a bare rel string is an
        open-folder command, a list is a further cascade. Overlong levels
        arrive pre-chunked into "first - last" sub-cascades."""
        for entry in _chunk_menu_entries(entries):
            if entry[0] == "-":
                menu.add_separator()
            elif isinstance(entry[1], str):
                menu.add_command(label=entry[0],
                                 command=lambda rr=entry[1]: self.open_job_folder(rr))
            else:
                sub = tk.Menu(menu, tearoff=0)
                self._fill_menu(sub, entry[1])
                menu.add_cascade(label=entry[0], menu=sub)

    def _render_items(self, frame, items, width):
        for i, it in enumerate(items):
            col = i % 3
            r = i // 3
            if it["kind"] == "dropdown":
                mb = ttk.Menubutton(frame, text="\u25BE " + it["label"],
                                    style="Card.TMenubutton", width=width)
                menu = tk.Menu(mb, tearoff=0)
                self._fill_menu(menu, it["menu"])
                mb["menu"] = menu
                mb.grid(row=r, column=col, sticky="ew", padx=5, pady=5)
            elif it["kind"] == "hover":
                self._hover_button(frame, it, width, r, col)
            else:
                btn = ttk.Button(
                    frame, text="{}  {}".format(it["icon"], it["label"]),
                    style="Card.TButton", width=width,
                    command=lambda rr=it["rel"]: self.open_job_folder(rr),
                )
                btn.grid(row=r, column=col, sticky="ew", padx=5, pady=5)
        frame.grid_columnconfigure(0, weight=1, uniform="btn")
        frame.grid_columnconfigure(1, weight=1, uniform="btn")
        frame.grid_columnconfigure(2, weight=1, uniform="btn")

    def _hover_button(self, frame, it, width, r, col):
        """Menubutton whose dropdown appears on hover; a click on the button
        itself opens its folder."""
        items = it.get("menu") or []
        if not items:
            ttk.Button(
                frame, text="{}  {}".format(it.get("icon", ""), it["label"]),
                style="Card.TButton", width=width,
                command=lambda rr=it["rel"]: self.open_job_folder(rr),
            ).grid(row=r, column=col, sticky="ew", padx=5, pady=5)
            return

        mb = ttk.Menubutton(frame, text="\u25BE " + it["label"],
                            style="Card.TMenubutton", width=width)
        menu = tk.Menu(mb, tearoff=0)
        for label, rel in items:
            menu.add_command(label=label,
                             command=lambda rr=rel: self.open_job_folder(rr))
        mb["menu"] = menu
        self._hover_menus.append(menu)

        def _post(_e):
            for m in self._hover_menus:
                try:
                    m.unpost()
                except tk.TclError:
                    pass
            if menu.index("end") is not None:
                menu.post(mb.winfo_rootx(), mb.winfo_rooty() + mb.winfo_height())

        def _leave(_e):
            x, y = mb.winfo_pointerxy()
            w = mb.winfo_containing(x, y)
            if w is None or not str(w).startswith(str(menu)):
                try:
                    menu.unpost()
                except tk.TclError:
                    pass

        def _click(_e):
            for m in self._hover_menus:
                try:
                    m.unpost()
                except tk.TclError:
                    pass
            self.open_job_folder(it["rel"])
            return "break"

        mb.bind("<Enter>", _post)
        mb.bind("<Leave>", _leave)
        mb.bind("<Button-1>", _click)
        mb.grid(row=r, column=col, sticky="ew", padx=5, pady=5)

    def _icon_for_type(self, typ):
        # Modern Windows (Segoe MDL2 Assets) glyphs instead of emoji.
        return default_icon_for_type(typ)

    def open_job_folder(self, rel):
        if not self.current_job:
            return
        path = os.path.join(JOB_ROOT, self.current_job, rel) if rel else os.path.join(JOB_ROOT, self.current_job)
        open_target(path)

    def open_current_job(self):
        if self.current_job:
            open_target(os.path.join(JOB_ROOT, self.current_job))

    def _add_recent(self, job):
        recent = self.settings.setdefault("recent", [])
        if job in recent:
            recent.remove(job)
        recent.insert(0, job)
        self.settings["recent"] = recent[:10]
        save_settings(self.settings)

    # ---------- favorites (shown directly in the sidebar) ----------

    def refresh_favorites(self):
        self.fav_list.delete(0, "end")
        favs = self.settings.get("favorites", [])
        for f in favs:
            self.fav_list.insert("end", "\u2605  {}".format(f))
        if not favs:
            self.fav_list.insert("end", "  (no favorites yet)")
        self._fit_sidebar_width()

    def _on_favorite_select(self, _=None):
        sel = self.fav_list.curselection()
        favs = self.settings.get("favorites", [])
        if not sel or sel[0] >= len(favs):
            return
        self._open_job(favs[sel[0]])

    def _favorite_open_folder(self):
        sel = self.fav_list.curselection()
        favs = self.settings.get("favorites", [])
        if not sel or sel[0] >= len(favs):
            return
        open_target(os.path.join(JOB_ROOT, favs[sel[0]]))

    def _job_context_menu(self, event):
        idx = self.job_list.nearest(event.y)
        jobs = self.filtered_jobs
        if idx < 0 or idx >= len(jobs):
            return
        name = jobs[idx]
        self.job_list.selection_clear(0, "end")
        self.job_list.selection_set(idx)
        menu = tk.Menu(self, tearoff=0)
        menu.add_command(label="Open in Explorer",
                         command=lambda: open_target(os.path.join(JOB_ROOT, name)))
        menu.add_command(label="Open in Navigator", command=lambda: self._open_job(name))
        menu.add_separator()
        menu.add_command(label="Add to Favorites",
                         command=lambda: self._add_favorite(name))
        menu.tk_popup(event.x_root, event.y_root)

    def _favorite_context_menu(self, event):
        idx = self.fav_list.nearest(event.y)
        favs = self.settings.get("favorites", [])
        if idx < 0 or idx >= len(favs):
            return
        name = favs[idx]
        self.fav_list.selection_clear(0, "end")
        self.fav_list.selection_set(idx)
        menu = tk.Menu(self, tearoff=0)
        menu.add_command(label="Open", command=lambda: self._open_job(name))
        menu.add_separator()
        menu.add_command(label="Remove from favorites",
                         command=lambda: self._remove_favorite(name))
        menu.tk_popup(event.x_root, event.y_root)

    def _remove_favorite(self, name):
        favs = self.settings.get("favorites", [])
        if name in favs:
            favs.remove(name)
            self.settings["favorites"] = favs
            save_settings(self.settings)
        self.refresh_favorites()
        self.status_var.set("Removed '{}' from favorites.".format(name))

    def _add_favorite(self, name):
        favs = self.settings.setdefault("favorites", [])
        if name in favs:
            return False
        favs.append(name)
        save_settings(self.settings)
        self.refresh_favorites()
        self.status_var.set("Added '{}' to favorites.".format(name))
        return True

    def add_favorite(self):
        if not self.current_job:
            messagebox.showinfo(
                APP_NAME,
                "Select a job from the list first, then click Add to Favorites.",
                parent=self.winfo_toplevel())
            return
        if not self._add_favorite(self.current_job):
            messagebox.showinfo(
                APP_NAME,
                "'{}' is already a favorite.".format(self.current_job),
                parent=self.winfo_toplevel())
            return
        messagebox.showinfo(
            APP_NAME,
            "Added '{}' to Favorites (shown in the sidebar).".format(self.current_job),
            parent=self.winfo_toplevel())

    def _check_admin(self, parent):
        """Admin username/password gate. Returns True on success - used
        before any write to the shared admin links."""
        creds = load_admin_credentials()
        if not creds:
            messagebox.showerror(
                APP_NAME,
                "Could not read the admin config (DEA_Config.xlsx).\nAdmin links are locked.",
                parent=parent)
            return False
        user = simpledialog.askstring(APP_NAME, "Admin username:", parent=parent)
        if user is None:
            return False
        pwd = simpledialog.askstring(APP_NAME, "Admin password:", show="*", parent=parent)
        if pwd is None:
            return False
        if (user.strip(), pwd) != creds:
            messagebox.showerror(APP_NAME, "Invalid admin credentials.", parent=parent)
            return False
        return True

    def _edit_link_dialog(self, link, parent):
        """Modal name + target editor. Returns (name, target), or None if
        cancelled / left incomplete."""
        result = {}

        def save():
            name = name_var.get().strip()
            target = target_var.get().strip()
            if not name or not target:
                messagebox.showwarning(APP_NAME, "Both a name and a URL or path are needed.",
                                       parent=win)
                return
            result.update(name=name, target=target)
            win.destroy()

        win = tk.Toplevel(parent)
        win.title("Edit Link")
        win.resizable(False, False)
        win.transient(parent)
        win.grab_set()
        ttk.Label(win, text="Name:").grid(row=0, column=0, sticky="w", padx=14, pady=(12, 2))
        name_var = tk.StringVar(value=link.get("name", ""))
        ttk.Entry(win, textvariable=name_var, width=48).grid(row=1, column=0, padx=14, sticky="ew")
        ttk.Label(win, text="URL or path:").grid(row=2, column=0, sticky="w", padx=14, pady=(8, 2))
        target_var = tk.StringVar(value=link.get("target", ""))
        ttk.Entry(win, textvariable=target_var, width=48).grid(row=3, column=0, padx=14, sticky="ew")
        btns = ttk.Frame(win)
        btns.grid(row=4, column=0, sticky="e", padx=14, pady=12)
        ttk.Button(btns, text="Save", command=save).pack(side="left")
        ttk.Button(btns, text="Cancel", command=win.destroy).pack(side="left", padx=(6, 0))
        win.bind("<Escape>", lambda e: win.destroy())
        win.bind("<Return>", lambda e: save())
        parent.wait_window(win)
        if result:
            return result["name"], result["target"]
        return None

    def _render_links_manager(self, host, row):
        """The link manager embedded at the bottom of the Quick Links tab
        (add/rename/delete/reorder/icons/admin links) - no popup. Every
        edit repaints only this tab (see _refresh_quick_links) so the
        edited row stays selected."""
        top = self.winfo_toplevel()
        toggle = tk.Label(host, text="\u25BE  MANAGE LINKS",
                          font=("Segoe UI", 9, "bold"),
                          bg="#f5f6f8", fg="#20252b", cursor="hand2")
        toggle.grid(row=row, column=0, columnspan=3, sticky="w", padx=8, pady=(14, 0))
        body = ttk.Frame(host, style="App.TFrame")
        body.grid(row=row + 1, column=0, columnspan=3, sticky="ew", padx=2)
        body.grid_columnconfigure(0, weight=1)

        tree_frame = ttk.Frame(body)
        tree_frame.grid(row=0, column=0, sticky="ew", pady=(4, 0))
        tree_frame.grid_rowconfigure(0, weight=1)
        tree_frame.grid_columnconfigure(0, weight=1)
        tree = ttk.Treeview(tree_frame, columns=("type", "target"), show="tree headings",
                            style="Nav.Treeview", height=8)
        tree.heading("#0", text="Name")
        tree.heading("type", text="Type")
        tree.heading("target", text="Target")
        for _col, _w in (("#0", 180), ("type", 70), ("target", 340)):
            tree.column(_col, width=_w, stretch=False)
        tree_vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        tree_hsb = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=tree_vsb.set, xscrollcommand=tree_hsb.set)
        tree.grid(row=0, column=0, sticky="ew")
        tree_vsb.grid(row=0, column=1, sticky="ns")
        tree_hsb.grid(row=1, column=0, sticky="ew")
        tree.bind("<Double-Button-1>", lambda e: edit_link())
        self._links_tree = tree

        hint = tk.StringVar(value=self._links_notice)
        self._links_notice = ""
        ttk.Label(body, textvariable=hint, foreground="#68727d",
                  font=("Segoe UI", 9)).grid(row=1, column=0, sticky="w")

        def all_links():
            return get_admin_links() + self.settings.get("links", [])

        def reload_tree():
            for x in tree.get_children():
                tree.delete(x)
            used_colors = set()
            for idx, link in enumerate(all_links()):
                admin = idx < len(get_admin_links())
                shown = "{} {}".format(resolve_link_icon(link), link["name"])
                color = link.get("color")
                tags = ()
                if color:
                    tags = ("linkcolor_" + str(color),)
                    used_colors.add(str(color))
                tree.insert("", "end", iid=str(idx),
                            values=("Admin" if admin else link["type"], link["target"]),
                            text=("{} (Admin)".format(shown) if admin else shown),
                            tags=tags)
            for color in used_colors:
                tree.tag_configure("linkcolor_" + color, foreground=color)

        def add_link():
            name = simpledialog.askstring(APP_NAME, "Link name:", parent=top)
            if not name:
                return
            target = simpledialog.askstring(APP_NAME, "URL or path:", parent=top)
            if not target:
                return
            self.settings.setdefault("links", []).append(
                {"name": name, "type": detect_link_type(target), "target": target})
            save_settings(self.settings)
            self._refresh_quick_links(select=len(all_links()) - 1)

        def delete_link():
            sel = tree.selection()
            if not sel:
                return
            idx = int(sel[0])
            if idx < len(get_admin_links()):
                messagebox.showwarning(
                    APP_NAME, "Admin links are fixed and cannot be deleted.", parent=top)
                return
            del self.settings["links"][idx - len(get_admin_links())]
            save_settings(self.settings)
            self._refresh_quick_links(select=min(idx, len(all_links()) - 1))

        def edit_link(_=None):
            sel = tree.selection()
            if not sel:
                return
            idx = int(sel[0])
            if idx < len(get_admin_links()):
                if not self._check_admin(top):
                    return
                admin_list = get_admin_links()
                result = self._edit_link_dialog(admin_list[idx], top)
                if not result:
                    return
                admin_list[idx]["name"], admin_list[idx]["target"] = result
                admin_list[idx]["type"] = detect_link_type(result[1])
                try:
                    _write_links_json(SHARED_ADMIN_LINKS, admin_list)
                except Exception as e:
                    messagebox.showerror(APP_NAME, "Could not save admin links:\n{}".format(e),
                                         parent=top)
                    return
                invalidate_admin_links_cache()
                self._refresh_quick_links(select=idx)
                return
            link = self.settings["links"][idx - len(get_admin_links())]
            result = self._edit_link_dialog(link, top)
            if not result:
                return
            link["name"], link["target"] = result
            link["type"] = detect_link_type(result[1])
            save_settings(self.settings)
            self._refresh_quick_links(select=idx)

        def set_link_icon(_=None):
            sel = tree.selection()
            if not sel:
                return
            idx = int(sel[0])
            if idx < len(get_admin_links()):
                messagebox.showwarning(
                    APP_NAME, "Admin links are fixed here \u2014 change icons via Admin Links....", parent=top)
                return
            link = self.settings["links"][idx - len(get_admin_links())]
            picked = pick_link_icon(top, link.get("icon"), link.get("color"))
            if not picked:
                return
            glyph, color = picked
            if glyph:
                link["icon"] = glyph
            link["color"] = color
            save_settings(self.settings)
            self._refresh_quick_links(select=idx)

        def move_link(step):
            sel = tree.selection()
            if not sel:
                return
            idx = int(sel[0])
            admin_count = len(get_admin_links())
            if idx < admin_count:
                hint.set("Admin links are fixed here \u2014 reorder them via Admin Links....")
                return
            user_links = self.settings.setdefault("links", [])
            pos, new_pos = idx - admin_count, idx - admin_count + step
            if new_pos < 0 or new_pos >= len(user_links):
                return
            user_links[pos], user_links[new_pos] = user_links[new_pos], user_links[pos]
            save_settings(self.settings)
            self._refresh_quick_links(select=idx + step)

        def add_dropped_links(data):
            added = 0
            for raw in parse_drop_paths(data):
                unc = to_unc(raw)
                if not unc:
                    continue
                name = os.path.basename(unc.rstrip("\\")) or unc
                self.settings.setdefault("links", []).append(
                    {"name": name, "type": detect_link_type(unc), "target": unc})
                added += 1
            if added:
                save_settings(self.settings)
                self._links_notice = "Added {} link(s) \u2014 UNC path generated automatically.".format(added)
                self._refresh_quick_links(select=len(all_links()) - 1)

        if HAS_DND:
            try:
                tree.drop_target_register(DND_FILES)
                tree.dnd_bind("<<Drop>>", lambda e: add_dropped_links(e.data))
                if not hint.get():
                    hint.set("Drag files or folders onto the list \u2014 UNC path is generated for you.")
            except Exception:
                pass

        buttons = ttk.Frame(body)
        buttons.grid(row=2, column=0, sticky="w", pady=(4, 0))
        ttk.Button(buttons, text="+ Add Link", command=add_link).pack(side="left")
        ttk.Button(buttons, text="Delete", command=delete_link).pack(side="left", padx=6)
        ttk.Button(buttons, text="\u25B2 Up", command=lambda: move_link(-1)).pack(side="left")
        ttk.Button(buttons, text="\u25BC Down", command=lambda: move_link(1)).pack(side="left", padx=6)
        ttk.Button(buttons, text="Icon...", command=set_link_icon).pack(side="left")
        ttk.Button(buttons, text="Admin Links...",
                   command=lambda: self._edit_admin_links(top)).pack(side="left", padx=6)
        toggle.bind("<Button-1>", lambda e, h=toggle, f=body:
                    self._toggle_section("MANAGE LINKS", h, f, "MANAGE LINKS"))
        if "MANAGE LINKS" in self._collapsed:
            body.grid_remove()
            toggle.configure(text="\u25B8  MANAGE LINKS")
        reload_tree()

    def _edit_admin_links(self, parent=None):
        """The admin edits the shared admin links here (saved to admin_links.json
        next to the shared DEA_Config.xlsx so every install sees them). Unlocked
        only with the credentials stored in the shared DEA_Config.xlsx."""
        top = parent or self.winfo_toplevel()
        if not self._check_admin(top):
            return

        path = SHARED_ADMIN_LINKS
        win = tk.Toplevel(self)
        win.title("NaviTool 2.0 - Admin Quick Links")
        win.geometry("700x450")
        win.minsize(620, 400)
        win.transient(parent or self.winfo_toplevel())
        win.grab_set()
        links = [dict(l) for l in get_admin_links()]

        tree_frame = ttk.Frame(win)
        tree_frame.pack(fill="both", expand=True, padx=12, pady=(12, 4))
        tree_frame.grid_rowconfigure(0, weight=1)
        tree_frame.grid_columnconfigure(0, weight=1)
        tree = ttk.Treeview(tree_frame, columns=("type", "target"), show="tree headings",
                               style="Nav.Treeview")
        tree.heading("#0", text="Name")
        tree.heading("type", text="Type")
        tree.heading("target", text="Target")
        for _col, _w in (("#0", 200), ("type", 70), ("target", 340)):
            tree.column(_col, width=_w, stretch=False)
        tree_vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        tree_hsb = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=tree_vsb.set, xscrollcommand=tree_hsb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        tree_vsb.grid(row=0, column=1, sticky="ns")
        tree_hsb.grid(row=1, column=0, sticky="ew")
        tree.bind("<Double-Button-1>", lambda e: edit())

        def reload_tree():
            for x in tree.get_children():
                tree.delete(x)
            used_colors = set()
            for i, l in enumerate(links):
                shown = "{} {}".format(resolve_link_icon(l), l["name"])
                color = l.get("color")
                tags = ()
                if color:
                    tags = ("linkcolor_" + str(color),)
                    used_colors.add(str(color))
                tree.insert("", "end", iid=str(i), values=(l["type"], l["target"]),
                            text=shown, tags=tags)
            for color in used_colors:
                tree.tag_configure("linkcolor_" + color, foreground=color)

        def add():
            name = simpledialog.askstring(APP_NAME, "Admin link name:", parent=win)
            if not name:
                return
            target = simpledialog.askstring(APP_NAME, "URL or path:", parent=win)
            if not target:
                return
            links.append({"name": name, "type": detect_link_type(target), "target": target})
            reload_tree()

        def delete():
            sel = tree.selection()
            if not sel:
                return
            del links[int(sel[0])]
            reload_tree()

        def edit(_=None):
            sel = tree.selection()
            if not sel:
                return
            i = int(sel[0])
            result = self._edit_link_dialog(links[i], win)
            if not result:
                return
            links[i]["name"], links[i]["target"] = result
            links[i]["type"] = detect_link_type(result[1])
            reload_tree()
            tree.selection_set(str(i))
            tree.see(str(i))

        def set_icon(_=None):
            sel = tree.selection()
            if not sel:
                return
            i = int(sel[0])
            picked = pick_link_icon(win, links[i].get("icon"), links[i].get("color"))
            if not picked:
                return
            glyph, color = picked
            if glyph:
                links[i]["icon"] = glyph
            links[i]["color"] = color
            reload_tree()
            tree.selection_set(str(i))
            tree.see(str(i))

        def move(step):
            sel = tree.selection()
            if not sel:
                return
            pos, new_pos = int(sel[0]), int(sel[0]) + step
            if new_pos < 0 or new_pos >= len(links):
                return
            links[pos], links[new_pos] = links[new_pos], links[pos]
            reload_tree()
            tree.selection_set(str(new_pos))
            tree.see(str(new_pos))

        def save():
            try:
                _write_links_json(path, links)
            except Exception as e:
                messagebox.showerror(APP_NAME, "Could not save admin links:\n{}".format(e), parent=win)
                return
            invalidate_admin_links_cache()
            messagebox.showinfo(
                APP_NAME,
                "Admin links saved to the shared config:\n{}\n\nAll installs share this file.".format(path),
                parent=win)
            win.destroy()
            self._refresh_quick_links()

        hint = tk.StringVar(value="")

        def add_dropped_links(data):
            added = 0
            for raw in parse_drop_paths(data):
                unc = to_unc(raw)
                if not unc:
                    continue
                name = os.path.basename(unc.rstrip("\\")) or unc
                links.append({"name": name, "type": detect_link_type(unc), "target": unc})
                added += 1
            if added:
                reload_tree()
                hint.set("Added {} link(s) \u2014 UNC path generated automatically.".format(added))

        if HAS_DND:
            try:
                win.drop_target_register(DND_FILES)
                win.dnd_bind("<<Drop>>", lambda e: add_dropped_links(e.data))
                hint.set("Drag files or folders here from Explorer \u2014 UNC path is generated for you.")
            except Exception:
                pass

        ttk.Label(win, text="Shared by all installs. Saved to (next to DEA_Config.xlsx):\n" + path,
                  foreground="#68727d", font=("Segoe UI", 9)).pack(anchor="w", padx=12)
        ttk.Label(win, textvariable=hint, foreground="#68727d",
                  font=("Segoe UI", 9)).pack(anchor="w", padx=12)
        btns = ttk.Frame(win)
        btns.pack(fill="x", padx=12, pady=(0, 2))
        ttk.Button(btns, text="+ Add", command=add).pack(side="left")
        ttk.Button(btns, text="Delete", command=delete).pack(side="left", padx=6)
        ttk.Button(btns, text="\u25B2 Up", command=lambda: move(-1)).pack(side="left")
        ttk.Button(btns, text="\u25BC Down", command=lambda: move(1)).pack(side="left", padx=6)
        ttk.Button(btns, text="Icon...", command=set_icon).pack(side="left")
        btns2 = ttk.Frame(win)
        btns2.pack(fill="x", padx=12, pady=(0, 12))
        ttk.Button(btns2, text="Save", command=save).pack(side="right")
        ttk.Button(btns2, text="Cancel", command=win.destroy).pack(side="right", padx=6)
        reload_tree()

def resource_path(name):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)